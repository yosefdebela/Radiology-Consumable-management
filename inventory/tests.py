from decimal import Decimal

from django.contrib.auth import get_user_model
from django.core.cache import cache
from django.db import IntegrityError, transaction
from django.test import Client, TestCase, override_settings
from django.urls import reverse

from .models import (CatalogItem, Consumable, ConsumableAlternative, ConsumableCategory, ConsumableSet,
                     ConsumableSetItem,
                     Examination, ExaminationConsumable, Modality, Unit)


def make_user(name, edit=False, delete=False, admin=False, change=False):
    from .models import UserProfile
    u = get_user_model().objects.create_user(name, password='pw-Strong-123', is_superuser=admin, is_staff=admin)
    UserProfile.objects.create(user=u, can_edit=edit or admin, can_change=change or admin,
                               can_delete=delete or admin)
    return u


@override_settings(PASSWORD_HASHERS=['django.contrib.auth.hashers.MD5PasswordHasher'])   # fast tests only
class BaseCase(TestCase):
    def setUp(self):
        cache.clear()
        self.admin = make_user('boss', admin=True)
        self.client.force_login(self.admin)
        self.ct = Modality.objects.get(slug='ct')
        self.mri = Modality.objects.get(slug='mri')
        self.ml, self.pair, self.piece, self.cm = (Unit.objects.get(name=n) for n in ('ml', 'pair', 'piece', 'cm'))
        self.exam = Examination.objects.create(name='CT Abdomen/Pelvis with contrast', code='X1', modality=self.ct)
        self.exam2 = Examination.objects.create(name='MRI Brain', code='X2', modality=self.mri)
        self.contrast = Consumable.objects.create(name='Contrast', default_unit=self.ml)
        self.cannula = Consumable.objects.create(name='IV Cannula', default_unit=self.piece)


class PagesRender(BaseCase):
    def test_all_pages_render(self):
        urls = [reverse('dashboard'), reverse('exam_list'), reverse('modality_exam_list', args=['ct']),
                reverse('exam_create'), reverse('exam_detail', args=[self.exam.pk]),
                reverse('exam_edit', args=[self.exam.pk]), reverse('consumable_list'),
                reverse('consumable_create'), reverse('consumable_edit', args=[self.contrast.pk]),
                reverse('category_list'), reverse('unit_list'), reverse('modality_list'),
                reverse('modality_edit', args=[self.ct.pk]), reverse('bulk_apply'), reverse('export_excel')]
        for u in urls:
            self.assertEqual(self.client.get(u).status_code, 200, u)


class Dashboard(BaseCase):
    def test_counts(self):
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast,
                                             quantity=3, unit=self.pair)
        r = self.client.get(reverse('dashboard'))
        self.assertEqual(r.context['total_exams'], 2)
        self.assertEqual(r.context['missing_count'], 1)
        self.assertEqual(r.context['total_consumables'], 2 + 0)
        by = {m['obj'].slug: m['total'] for m in r.context['modalities']}
        self.assertEqual(by['ct'], 1)
        self.assertEqual(by['mri'], 1)
        self.assertEqual(by['ultrasound'], 0)

    def test_inactive_excluded(self):
        self.exam2.is_active = False
        self.exam2.save()
        self.assertEqual(self.client.get(reverse('dashboard')).context['total_exams'], 1)


class ExamConsumables(BaseCase):
    def test_add_update_remove(self):
        r = self.client.post(reverse('link_add', args=[self.exam.pk]), {
            'consumable': self.contrast.pk, 'quantity': '3', 'unit': self.pair.pk,
            'is_required': 'on', 'notes': 'three pair'})
        self.assertEqual(r.status_code, 302)
        link = ExaminationConsumable.objects.get()
        self.assertEqual((link.quantity, link.unit, link.is_required), (Decimal('3'), self.pair, True))

        # change quantity, unit, optional
        self.client.post(reverse('link_update', args=[link.pk]), {
            f'l{link.pk}-quantity': '15', f'l{link.pk}-unit': self.cm.pk,
            f'l{link.pk}-is_required': '', f'l{link.pk}-notes': 'plaster'})
        link.refresh_from_db()
        self.assertEqual((link.quantity, link.unit, link.is_required), (Decimal('15'), self.cm, False))

        self.client.post(reverse('link_remove', args=[link.pk]))
        self.assertFalse(ExaminationConsumable.objects.exists())

    def test_duplicate_and_bad_quantity_rejected(self):
        url = reverse('link_add', args=[self.exam.pk])
        data = {'consumable': self.contrast.pk, 'quantity': '1', 'unit': self.ml.pk, 'is_required': 'on'}
        self.client.post(url, data)
        self.client.post(url, data)
        self.client.post(url, {**data, 'consumable': self.cannula.pk, 'quantity': '0'})
        self.assertEqual(ExaminationConsumable.objects.count(), 1)

    def test_inactive_consumable_not_offered(self):
        self.cannula.is_active = False
        self.cannula.save()
        r = self.client.get(reverse('exam_detail', args=[self.exam.pk]))
        self.assertNotIn(self.cannula, r.context['new_forms'][0].fields['consumable'].queryset)

    def test_copy(self):
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast,
                                             quantity=2, unit=self.ml)
        self.client.post(reverse('exam_copy', args=[self.exam2.pk]), {'source': self.exam.pk})
        self.assertEqual(self.exam2.consumable_links.get().quantity, Decimal('2'))


class ConsumableRules(BaseCase):
    def test_delete_unused_removes(self):
        self.client.post(reverse('consumable_delete', args=[self.cannula.pk]))
        self.assertFalse(Consumable.objects.filter(pk=self.cannula.pk).exists())

    def test_delete_used_deactivates(self):
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast,
                                             quantity=1, unit=self.ml)
        self.client.post(reverse('consumable_delete', args=[self.contrast.pk]))
        self.contrast.refresh_from_db()
        self.assertFalse(self.contrast.is_active)
        self.assertEqual(ExaminationConsumable.objects.count(), 1)

    def test_db_level_protect(self):
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast,
                                             quantity=1, unit=self.ml)
        from django.db.models import ProtectedError
        with self.assertRaises(ProtectedError):
            self.contrast.delete()

    def test_search_and_category_filter(self):
        cat = ConsumableCategory.objects.create(name='Test cat')
        self.contrast.category = cat
        self.contrast.save()
        r = self.client.get(reverse('consumable_list'), {'q': 'cannula'})
        self.assertEqual([c.name for c in r.context['page']], ['IV Cannula'])
        r = self.client.get(reverse('consumable_list'), {'category': cat.pk})
        self.assertEqual([c.name for c in r.context['page']], ['Contrast'])

    def test_create_edit_toggle(self):
        self.client.post(reverse('consumable_create'), {'name': 'Gel', 'default_unit': self.ml.pk,
                                                       'is_active': 'on'})
        g = Consumable.objects.get(name='Gel')
        self.client.post(reverse('consumable_edit', args=[g.pk]), {'name': 'US Gel', 'default_unit': self.ml.pk})
        g.refresh_from_db()
        self.assertEqual(g.name, 'US Gel')
        self.assertFalse(g.is_active)  # unchecked box = inactive
        self.client.post(reverse('consumable_toggle', args=[g.pk]))
        g.refresh_from_db()
        self.assertTrue(g.is_active)


class ExamRules(BaseCase):
    def test_search_filter_toggle_create(self):
        r = self.client.get(reverse('exam_list'), {'q': 'brain'})
        self.assertEqual([e.name for e in r.context['page']], ['MRI Brain'])
        r = self.client.get(reverse('modality_exam_list', args=['ct']))
        self.assertEqual([e.name for e in r.context['page']], ['CT Abdomen/Pelvis with contrast'])
        self.client.post(reverse('exam_toggle', args=[self.exam.pk]))
        self.exam.refresh_from_db()
        self.assertFalse(self.exam.is_active)
        self.assertEqual(self.client.get(reverse('exam_list')).context['page'].paginator.count, 1)
        self.client.post(reverse('exam_create'), {'name': 'CT Neck', 'modality': self.ct.pk, 'is_active': 'on'})
        self.assertTrue(Examination.objects.filter(name='CT Neck').exists())

    def test_missing_filter(self):
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast,
                                             quantity=1, unit=self.ml)
        r = self.client.get(reverse('exam_list'), {'missing': '1'})
        self.assertEqual([e.name for e in r.context['page']], ['MRI Brain'])


class NewFeatures(BaseCase):
    def test_new_modality_from_exam_form(self):
        r = self.client.post(reverse('exam_create'), {'name': 'DEXA Spine', 'modality': '', 'new_modality': 'Bone Densitometry',
                                                     'is_active': 'on'})
        self.assertEqual(r.status_code, 302)
        m = Modality.objects.get(name='Bone Densitometry')
        self.assertEqual(m.slug, 'bone-densitometry')
        self.assertTrue(m.color.startswith('#'))
        self.assertEqual(Examination.objects.get(name='DEXA Spine').modality, m)
        self.assertContains(self.client.get(reverse('dashboard')), 'Bone Densitometry')
        self.assertContains(self.client.get(reverse('dashboard')), '.mod-bone-densitometry')

    def test_new_modality_name_reuses_existing(self):
        before = Modality.objects.count()
        self.client.post(reverse('exam_create'), {'name': 'CT Neck', 'new_modality': ' ct   scan ', 'is_active': 'on'})
        self.assertEqual(Modality.objects.count(), before)
        self.assertEqual(Examination.objects.get(name='CT Neck').modality, self.ct)

    def test_modality_required(self):
        r = self.client.post(reverse('exam_create'), {'name': 'No modality', 'is_active': 'on'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(Examination.objects.filter(name='No modality').exists())

    def test_create_with_copy(self):
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast, quantity=3, unit=self.pair)
        self.client.post(reverse('exam_create'), {'name': 'CT Chest with contrast', 'modality': self.ct.pk,
                                                 'copy_from': self.exam.pk, 'is_active': 'on'})
        new = Examination.objects.get(name='CT Chest with contrast')
        self.assertEqual(new.consumable_links.get().quantity, Decimal('3'))

    def test_delete_only_when_inactive(self):
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast, quantity=1, unit=self.ml)
        self.client.post(reverse('exam_delete', args=[self.exam.pk]))
        self.assertTrue(Examination.objects.filter(pk=self.exam.pk).exists())
        self.client.post(reverse('exam_toggle', args=[self.exam.pk]))
        self.client.post(reverse('exam_delete', args=[self.exam.pk]))
        self.assertFalse(Examination.objects.filter(pk=self.exam.pk).exists())
        self.assertTrue(Consumable.objects.filter(pk=self.contrast.pk).exists())
        self.assertEqual(self.client.get(self.client.get(reverse('exam_list')).request['PATH_INFO']).status_code, 200)

    def test_bulk_apply_rows_and_copy(self):
        e3 = Examination.objects.create(name='MRI Knee', modality=self.mri)
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast, quantity=3, unit=self.pair)
        # exam2 already has the cannula at qty 1 (should be skipped), e3 has nothing
        ExaminationConsumable.objects.create(examination=self.exam2, consumable=self.cannula, quantity=1, unit=self.piece)
        data = {'exams': [self.exam2.pk, e3.pk], 'source': self.exam.pk, 'conflict': 'skip',
                'row_consumable': [self.cannula.pk, ''], 'row_quantity': ['2', '1'], 'row_unit': [self.piece.pk, ''],
                'row_required': ['0', '1'], 'row_notes': ['x', '']}
        r = self.client.post(reverse('bulk_apply'), data)
        self.assertEqual(r.status_code, 302)
        self.assertEqual(self.exam2.consumable_links.count(), 2)      # contrast copied, cannula untouched
        self.assertEqual(self.exam2.consumable_links.get(consumable=self.cannula).quantity, Decimal('1'))
        self.assertEqual(e3.consumable_links.count(), 2)
        self.assertEqual(e3.consumable_links.get(consumable=self.cannula).quantity, Decimal('2'))
        self.assertFalse(e3.consumable_links.get(consumable=self.cannula).is_required)

        data['conflict'] = 'overwrite'
        self.client.post(reverse('bulk_apply'), data)
        self.assertEqual(self.exam2.consumable_links.get(consumable=self.cannula).quantity, Decimal('2'))
        self.assertEqual(ExaminationConsumable.objects.filter(consumable=self.cannula).count(), 2)

    def test_bulk_apply_validation(self):
        base = {'row_consumable': [self.cannula.pk], 'row_quantity': ['0'], 'row_unit': [self.piece.pk],
                'row_required': ['1'], 'row_notes': ['']}
        for data in ({**base, 'exams': [self.exam.pk]},          # bad quantity
                     {**base, 'row_quantity': ['1']},             # no examinations selected
                     {'exams': [self.exam.pk]}):                  # nothing to apply
            r = self.client.post(reverse('bulk_apply'), data)
            self.assertEqual(r.status_code, 200)
        self.assertFalse(ExaminationConsumable.objects.exists())

    def test_bulk_page_preselect(self):
        r = self.client.get(reverse('bulk_apply'), {'exam': [self.exam.pk, self.exam2.pk]})
        self.assertEqual(r.context['preselected'], {self.exam.pk, self.exam2.pk})

    def test_export_workbook(self):
        from io import BytesIO
        from openpyxl import load_workbook
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast, quantity=Decimal('2.5'),
                                             unit=self.ml, notes='=1+1')
        r = self.client.get(reverse('export_excel'), {'download': '1'})
        self.assertEqual(r.status_code, 200)
        self.assertIn('.xlsx', r['Content-Disposition'])
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(wb.sheetnames, ['Modalities', 'Examination consumables', 'Consumables', 'Saved sets'])
        rows = list(wb['Examination consumables'].iter_rows(values_only=True))
        self.assertEqual(rows[0][2], 'Examination')
        by_exam = {r[2]: r for r in rows[1:]}
        self.assertEqual(by_exam['CT Abdomen/Pelvis with contrast'][5:10], ('Contrast', None, 2.5, 'ml', 'Required'))
        self.assertEqual(by_exam['CT Abdomen/Pelvis with contrast'][10], '=1+1')   # kept as text, not a formula
        self.assertIn('MRI Brain', by_exam)                                          # exam without consumables listed
        mods = {r[0]: r for r in wb['Modalities'].iter_rows(min_row=2, values_only=True)}
        self.assertEqual(mods['CT Scan'][1:], (1, 1, 0))

        r = self.client.get(reverse('export_excel'), {'download': '1', 'modality': 'mri', 'empty': '0'})
        wb = load_workbook(BytesIO(r.content))
        self.assertEqual(len(list(wb['Examination consumables'].iter_rows(min_row=2))), 0)
        self.assertEqual(self.client.get(reverse('export_excel')).status_code, 200)

    def test_modality_management(self):
        self.client.post(reverse('modality_list'), {'name': 'Mammography', 'color': '#112233', 'sort_order': 9})
        m = Modality.objects.get(name='Mammography')
        self.assertEqual(m.color, '#112233')
        self.client.post(reverse('modality_edit', args=[m.pk]), {'name': 'Mammography 2', 'color': '#445566', 'sort_order': 7})
        m.refresh_from_db()
        self.assertEqual((m.name, m.slug, m.color), ('Mammography 2', 'mammography', '#445566'))
        self.client.post(reverse('modality_delete', args=[self.ct.pk]))   # has exams -> kept
        self.assertTrue(Modality.objects.filter(pk=self.ct.pk).exists())
        self.client.post(reverse('modality_delete', args=[m.pk]))
        self.assertFalse(Modality.objects.filter(pk=m.pk).exists())


class WholeNumberUnits(BaseCase):
    def add(self, qty, unit):
        return self.client.post(reverse('link_add', args=[self.exam.pk]), {
            'consumable': self.contrast.pk, 'quantity': qty, 'unit': unit.pk, 'is_required': 'on'})

    def test_all_seeded_units_are_whole_number(self):
        self.assertFalse(Unit.objects.filter(allow_decimals=True).exists())

    def test_decimal_rejected_for_whole_number_unit(self):
        self.add('2.5', self.ml)
        self.assertFalse(ExaminationConsumable.objects.exists())
        self.add('3', self.ml)
        self.assertEqual(ExaminationConsumable.objects.get().quantity, Decimal('3'))

    def test_update_rejects_decimal(self):
        self.add('3', self.ml)
        link = ExaminationConsumable.objects.get()
        self.client.post(reverse('link_update', args=[link.pk]), {
            f'l{link.pk}-quantity': '1.5', f'l{link.pk}-unit': self.ml.pk, f'l{link.pk}-is_required': 'on',
            f'l{link.pk}-notes': ''})
        link.refresh_from_db()
        self.assertEqual(link.quantity, Decimal('3'))

    def test_unit_can_be_edited_to_allow_decimals(self):
        self.client.post(reverse('unit_edit', args=[self.ml.pk]), {'name': 'ml', 'allow_decimals': 'on', 'is_active': 'on'})
        self.ml.refresh_from_db()
        self.assertTrue(self.ml.allow_decimals)
        self.add('2.5', self.ml)
        self.assertEqual(ExaminationConsumable.objects.get().quantity, Decimal('2.50'))
        # cannot switch back while a decimal quantity exists
        self.client.post(reverse('unit_edit', args=[self.ml.pk]), {'name': 'ml', 'is_active': 'on'})
        self.ml.refresh_from_db()
        self.assertTrue(self.ml.allow_decimals)

    def test_bulk_rows_follow_unit_rule(self):
        data = {'exams': [self.exam.pk], 'row_consumable': [self.contrast.pk], 'row_quantity': ['1.5'],
                'row_unit': [self.ml.pk], 'row_required': ['1'], 'row_notes': ['']}
        self.assertEqual(self.client.post(reverse('bulk_apply'), data).status_code, 200)
        self.assertFalse(ExaminationConsumable.objects.exists())
        data['row_quantity'] = ['2']
        self.assertEqual(self.client.post(reverse('bulk_apply'), data).status_code, 302)
        self.assertEqual(ExaminationConsumable.objects.get().quantity, Decimal('2'))

    def test_display_has_no_trailing_zeros(self):
        self.add('3', self.ml)
        html = self.client.get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        self.assertIn('value="3"', html)
        self.assertNotIn('3.00', html)

    def test_select_all_button_on_bulk_page(self):
        self.assertContains(self.client.get(reverse('bulk_apply')), 'Select all 2 examinations')



def client_for(user):
    c = Client()
    c.force_login(user)
    return c


class Accounts(BaseCase):
    def test_anonymous_goes_to_login(self):
        r = Client().get('/')
        self.assertEqual(r.status_code, 302)
        self.assertTrue(r['Location'].startswith('/login/'))

    def test_first_run_setup(self):
        get_user_model().objects.all().delete()
        anon = Client()
        self.assertEqual(anon.get('/').status_code, 302)
        self.assertTrue(anon.get('/').url.startswith('/setup/'))
        r = anon.post('/setup/', {'username': 'first', 'full_name': 'First Admin',
                                  'password1': 'Zebra-Lantern-42', 'password2': 'Zebra-Lantern-42'})
        self.assertEqual(r.status_code, 302)
        u = get_user_model().objects.get(username='first')
        self.assertTrue(u.is_superuser)
        self.assertEqual(anon.get('/').status_code, 200)               # signed in straight away
        # the setup page is closed once a user exists
        self.assertEqual(Client().get('/setup/').status_code, 302)
        self.assertEqual(Client().post('/setup/', {'username': 'x'}).status_code, 302)
        self.assertFalse(get_user_model().objects.filter(username='x').exists())

    def test_login_logout(self):
        c = Client()
        r = c.post('/login/', {'username': 'boss', 'password': 'pw-Strong-123'})
        self.assertEqual(r.status_code, 302)
        self.assertEqual(c.get('/').status_code, 200)
        self.assertEqual(c.post('/logout/').status_code, 302)
        self.assertEqual(c.get('/').status_code, 302)

    def test_login_lockout_after_five_failures(self):
        c = Client()
        for _ in range(5):
            self.assertEqual(c.post('/login/', {'username': 'boss', 'password': 'wrong'}).status_code, 200)
        r = c.post('/login/', {'username': 'boss', 'password': 'pw-Strong-123'})   # right password, but locked
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'Too many wrong passwords')
        self.assertEqual(c.get('/').status_code, 302)

    def test_profile_password_change(self):
        u = make_user('pat')
        c = client_for(u)
        self.assertEqual(c.get('/profile/').status_code, 200)
        r = c.post('/profile/', {'old_password': 'pw-Strong-123', 'new_password1': 'Orange-Window-77',
                                 'new_password2': 'Orange-Window-77'})
        self.assertEqual(r.status_code, 302)
        u.refresh_from_db()
        self.assertTrue(u.check_password('Orange-Window-77'))


class ViewOnlyUser(BaseCase):
    def setUp(self):
        super().setUp()
        self.viewer = make_user('viewer')
        self.c = client_for(self.viewer)
        self.link = ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast,
                                                         quantity=3, unit=self.ml)

    def test_can_view_and_export(self):
        for url in (reverse('dashboard'), reverse('exam_list'), reverse('exam_detail', args=[self.exam.pk]),
                    reverse('consumable_list'), reverse('export_excel'), reverse('profile')):
            self.assertEqual(self.c.get(url).status_code, 200, url)
        self.assertEqual(self.c.get(reverse('export_excel'), {'download': '1'}).status_code, 200)

    def test_buttons_are_hidden(self):
        self.assertNotContains(self.c.get(reverse('exam_list')), 'Add examination')
        self.assertNotContains(self.c.get(reverse('exam_list')), 'Deactivate')
        detail = self.c.get(reverse('exam_detail', args=[self.exam.pk]))
        for text in ('Add consumables', 'Remove', 'Edit examination', 'Deactivate'):
            self.assertNotContains(detail, text)
        self.assertContains(detail, 'Contrast')
        self.assertNotContains(self.c.get(reverse('consumable_list')), 'Add consumable')
        nav = self.c.get(reverse('dashboard')).content.decode()
        self.assertNotIn('Apply to many', nav)
        self.assertNotIn('>Users<', nav)

    def test_every_change_is_refused(self):
        before = ExaminationConsumable.objects.count()
        posts = [reverse('link_add', args=[self.exam.pk]), reverse('link_update', args=[self.link.pk]),
                 reverse('link_remove', args=[self.link.pk]), reverse('exam_toggle', args=[self.exam.pk]),
                 reverse('exam_delete', args=[self.exam.pk]), reverse('exam_copy', args=[self.exam.pk]),
                 reverse('consumable_toggle', args=[self.contrast.pk]),
                 reverse('consumable_delete', args=[self.cannula.pk]), reverse('category_delete', args=[1])
                 if ConsumableCategory.objects.filter(pk=1).exists() else reverse('unit_toggle', args=[self.ml.pk]),
                 reverse('modality_delete', args=[self.mri.pk]),
                 reverse('exam_create'), reverse('consumable_create'), reverse('bulk_apply')]
        for url in posts:
            self.c.post(url, {'name': 'x', 'modality': self.ct.pk})
        self.assertEqual(ExaminationConsumable.objects.count(), before)
        self.exam.refresh_from_db(); self.contrast.refresh_from_db()
        self.assertTrue(self.exam.is_active and self.contrast.is_active)
        self.assertTrue(Consumable.objects.filter(pk=self.cannula.pk).exists())
        self.assertTrue(Modality.objects.filter(pk=self.mri.pk).exists())
        self.assertFalse(Examination.objects.filter(name='x').exists())
        self.assertFalse(Consumable.objects.filter(name='x').exists())

    def test_form_pages_are_forbidden(self):
        for url in (reverse('exam_create'), reverse('exam_edit', args=[self.exam.pk]), reverse('consumable_create'),
                    reverse('bulk_apply'), reverse('unit_list'), reverse('modality_list'), reverse('catalog_import'),
                    reverse('user_list'), reverse('consumable_suggest')):
            self.assertEqual(self.c.get(url).status_code, 403, url)


class EditorWithoutDelete(BaseCase):
    def setUp(self):
        super().setUp()
        self.c = client_for(make_user('editor', edit=True, change=True))
        self.link = ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast,
                                                         quantity=3, unit=self.ml)

    def test_can_add_and_change_but_not_delete(self):
        self.c.post(reverse('link_add', args=[self.exam.pk]), {
            'consumable': self.cannula.pk, 'quantity': '2', 'unit': self.piece.pk, 'is_required': 'on'})
        self.assertEqual(self.exam.consumable_links.count(), 2)
        self.c.post(reverse('link_update', args=[self.link.pk]), {
            f'l{self.link.pk}-quantity': '9', f'l{self.link.pk}-unit': self.ml.pk, f'l{self.link.pk}-is_required': 'on',
            f'l{self.link.pk}-notes': ''})
        self.link.refresh_from_db()
        self.assertEqual(self.link.quantity, Decimal('9'))
        self.c.post(reverse('exam_toggle', args=[self.exam.pk]))
        self.exam.refresh_from_db()
        self.assertFalse(self.exam.is_active)
        # ...but nothing can be removed
        self.c.post(reverse('link_remove', args=[self.link.pk]))
        self.c.post(reverse('exam_delete', args=[self.exam.pk]))
        self.c.post(reverse('consumable_delete', args=[self.cannula.pk]))
        self.assertTrue(ExaminationConsumable.objects.filter(pk=self.link.pk).exists())
        self.assertTrue(Examination.objects.filter(pk=self.exam.pk).exists())
        self.assertTrue(Consumable.objects.filter(pk=self.cannula.pk).exists())

    def test_pages_show_save_but_no_remove_or_delete(self):
        html = self.c.get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        self.assertIn('Save', html)
        self.assertNotIn('Remove', html)
        self.assertNotIn('>Delete', self.c.get(reverse('consumable_list')).content.decode())

    def test_delete_permission_needs_edit(self):
        from .permissions import can_delete
        odd = make_user('odd')
        odd.inventory_profile.can_delete = True
        odd.inventory_profile.save()
        self.assertFalse(can_delete(odd))


class DeleteAllowedUser(BaseCase):
    def test_can_remove_and_delete(self):
        c = client_for(make_user('cleaner', edit=True, delete=True, change=True))
        link = ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast,
                                                    quantity=1, unit=self.ml)
        c.post(reverse('link_remove', args=[link.pk]))
        self.assertFalse(ExaminationConsumable.objects.exists())
        c.post(reverse('consumable_delete', args=[self.cannula.pk]))
        self.assertFalse(Consumable.objects.filter(pk=self.cannula.pk).exists())
        self.assertNotContains(c.get(reverse('user_list')), 'Reset password', status_code=403)


class UserManagement(BaseCase):
    def test_admin_creates_and_edits_users(self):
        r = self.client.post(reverse('user_create'), {
            'username': 'nurse1', 'full_name': 'Nurse One', 'password1': 'Maple-Harbor-55',
            'password2': 'Maple-Harbor-55', 'account_type': 'user'})
        self.assertEqual(r.status_code, 302)
        u = get_user_model().objects.get(username='nurse1')
        self.assertFalse(u.is_superuser)
        self.assertEqual((u.inventory_profile.can_edit, u.inventory_profile.can_change,
                          u.inventory_profile.can_delete), (False, False, False))   # view only by default
        self.client.post(reverse('user_edit', args=[u.pk]), {
            'full_name': 'Nurse One', 'account_type': 'user', 'can_edit': 'on', 'can_change': 'on',
            'can_delete': 'on', 'is_active': 'on'})
        u.inventory_profile.refresh_from_db()
        self.assertEqual((u.inventory_profile.can_edit, u.inventory_profile.can_change,
                          u.inventory_profile.can_delete), (True, True, True))
        self.client.post(reverse('user_password', args=[u.pk]),
                         {'new_password1': 'River-Candle-91', 'new_password2': 'River-Candle-91'})
        u.refresh_from_db()
        self.assertTrue(u.check_password('River-Candle-91'))

    def test_delete_without_edit_is_rejected(self):
        u = make_user('x1')
        r = self.client.post(reverse('user_edit', args=[u.pk]), {
            'full_name': '', 'account_type': 'user', 'can_delete': 'on', 'is_active': 'on'})
        self.assertEqual(r.status_code, 200)

    def test_weak_password_rejected(self):
        r = self.client.post(reverse('user_create'), {
            'username': 'weak', 'password1': '12345678', 'password2': '12345678', 'account_type': 'user'})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(get_user_model().objects.filter(username='weak').exists())

    def test_cannot_remove_last_admin_or_deactivate_self(self):
        r = self.client.post(reverse('user_edit', args=[self.admin.pk]), {
            'full_name': '', 'account_type': 'user', 'is_active': 'on'})
        self.assertEqual(r.status_code, 200)
        self.client.post(reverse('user_edit', args=[self.admin.pk]), {'full_name': '', 'account_type': 'admin'})
        self.admin.refresh_from_db()
        self.assertTrue(self.admin.is_superuser and self.admin.is_active)

    def test_deactivated_user_cannot_sign_in(self):
        u = make_user('gone', edit=True)
        self.client.post(reverse('user_edit', args=[u.pk]), {
            'full_name': '', 'account_type': 'user', 'can_edit': 'on'})          # is_active left off
        r = Client().post('/login/', {'username': 'gone', 'password': 'pw-Strong-123'})
        self.assertEqual(r.status_code, 200)


def excel_bytes(sheets):
    from io import BytesIO
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


class ConsumablesListImport(BaseCase):
    def upload(self, client, data, **extra):
        from django.core.files.uploadedfile import SimpleUploadedFile
        f = SimpleUploadedFile('list.xlsx', data, content_type='application/octet-stream')
        return client.post(reverse('catalog_import'), {'file': f, **extra}, follow=True)

    def test_import_and_reimport(self):
        data = excel_bytes({'Sheet1': [['Item', 'Type', 'UOM', 'Notes'],
                                       ['Gauze swab', 'Dressings', 'pcs', 'Sterile'],
                                       ['Gel', 'Ultrasound', 'ML', ''],
                                       ['  Syringe   50 ml ', 'Syringes', 'Piece', ''],
                                       [None, None, None, None],
                                       ['Weird', 'Other', 'furlong', '']]})
        r = self.upload(self.client, data)
        self.assertEqual(CatalogItem.objects.count(), 4)
        gauze = CatalogItem.objects.get(name='Gauze swab')
        self.assertEqual(gauze.default_unit.name, 'piece')                  # pcs -> piece
        self.assertEqual(gauze.category.name, 'Dressings')
        self.assertEqual(gauze.description, 'Sterile')
        self.assertEqual(CatalogItem.objects.get(name='Gel').default_unit, self.ml)   # ML matches ml
        self.assertTrue(CatalogItem.objects.filter(name='Syringe 50 ml').exists())     # spaces tidied
        self.assertTrue(Unit.objects.filter(name='furlong').exists())                  # unknown unit is created
        self.assertContains(r, 'furlong')                                              # ...and reported
        self.assertFalse(Consumable.objects.filter(name='Gauze swab').exists())         # suggestions only
        self.upload(self.client, data)
        self.assertEqual(CatalogItem.objects.count(), 4)

    def test_without_headers_in_single_sheet_uses_first_column(self):
        self.upload(self.client, excel_bytes({'S': [['Catheter'], ['Tourniquet']]}))
        self.assertEqual(set(CatalogItem.objects.values_list('name', flat=True)), {'Tourniquet'})
        # the first row is treated as the header, as documented in the template

    def test_multi_sheet_skips_unrecognised_sheets(self):
        data = excel_bytes({'Items': [['Name', 'Unit'], ['Tape', 'roll']], 'Notes': [['Hello', 'World'], ['a', 'b']]})
        r = self.upload(self.client, data)
        self.assertEqual(list(CatalogItem.objects.values_list('name', flat=True)), ['Tape'])
        self.assertContains(r, 'Skipped (no name column): Notes')

    def test_add_all_creates_consumables(self):
        self.upload(self.client, excel_bytes({'S': [['Name', 'Unit'], ['Tape', 'roll'], ['No unit', '']]}), add_all='1')
        self.assertEqual(Consumable.objects.get(name='Tape').default_unit.name, 'roll')
        self.assertEqual(Consumable.objects.get(name='No unit').default_unit.name, 'piece')

    def test_bad_files_are_rejected_politely(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        r = self.client.post(reverse('catalog_import'), {'file': SimpleUploadedFile('a.txt', b'hi')}, follow=True)
        self.assertContains(r, 'Excel .xlsx file')
        r = self.client.post(reverse('catalog_import'), {'file': SimpleUploadedFile('a.xlsx', b'not a workbook')}, follow=True)
        self.assertContains(r, 'could not be read')

    def test_viewer_cannot_import(self):
        c = client_for(make_user('v'))
        self.upload(c, excel_bytes({'S': [['Name'], ['Tape']]}))
        self.assertFalse(CatalogItem.objects.exists())

    def test_template_download(self):
        r = self.client.get(reverse('catalog_template'))
        self.assertEqual(r.status_code, 200)
        self.assertIn('.xlsx', r['Content-Disposition'])

    def test_suggestions(self):
        cat = ConsumableCategory.objects.get_or_create(name='Contrast media')[0]
        CatalogItem.objects.create(name='Contrast agent 300', category=cat, default_unit=self.ml, description='Iodinated')
        CatalogItem.objects.create(name='Saline flush', default_unit=self.piece)
        r = self.client.get(reverse('consumable_suggest'), {'q': 'contr'}).json()
        self.assertEqual([i['name'] for i in r['results']], ['Contrast agent 300'])
        item = r['results'][0]
        self.assertEqual((item['category_name'], item['unit_name'], item['unit_id'], item['exists']),
                         ('Contrast media', 'ml', self.ml.pk, False))
        self.assertEqual(item['description'], 'Iodinated')
        # name matching is in the middle too, and short input returns nothing
        self.assertEqual(len(self.client.get(reverse('consumable_suggest'), {'q': 'flush'}).json()['results']), 1)
        self.assertEqual(self.client.get(reverse('consumable_suggest'), {'q': 'c'}).json()['results'], [])
        # once the consumable exists, the suggestion is flagged
        Consumable.objects.create(name='contrast AGENT 300', default_unit=self.ml)
        again = self.client.get(reverse('consumable_suggest'), {'q': 'contr'}).json()['results'][0]
        self.assertTrue(again['exists'])

    def test_add_form_wires_up_suggestions_and_blocks_duplicates(self):
        page = self.client.get(reverse('consumable_create')).content.decode()
        self.assertIn('data-suggest="1"', page)
        self.assertIn('suggest.js', page)
        self.assertNotIn('data-suggest="1"', self.client.get(reverse('consumable_edit', args=[self.contrast.pk])).content.decode())
        r = self.client.post(reverse('consumable_create'), {'name': 'CONTRAST', 'default_unit': self.ml.pk, 'is_active': 'on'})
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'already exists')

    def test_command_and_search_is_html_safe(self):
        import tempfile, os
        from django.core.management import call_command
        with tempfile.NamedTemporaryFile(suffix='.xlsx', delete=False) as t:
            t.write(excel_bytes({'S': [['Name', 'Category'], ['<b>Bold</b> gloves', 'Gloves']]}))
        try:
            call_command('import_consumable_catalog', t.name, verbosity=0)
        finally:
            os.unlink(t.name)
        self.assertEqual(CatalogItem.objects.get().name, '<b>Bold</b> gloves')
        page = self.client.get(reverse('catalog_import')).content.decode()
        self.assertNotIn('<b>Bold</b>', page)       # escaped by the template



class ChangeIsAnAdminPrivilege(BaseCase):
    """Adding is one thing; changing what was already entered needs an admin, or an admin's permission."""

    def setUp(self):
        super().setUp()
        self.adder = client_for(make_user('adder', edit=True))
        self.link = ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast,
                                                         quantity=3, unit=self.ml)

    def test_adder_can_add_but_not_change(self):
        self.adder.post(reverse('link_add', args=[self.exam.pk]), {
            'consumable': self.cannula.pk, 'quantity': '2', 'unit': self.piece.pk, 'is_required': 'on'})
        self.assertEqual(self.exam.consumable_links.count(), 2)
        r = self.adder.post(reverse('consumable_create'), {'name': 'Fresh gel', 'default_unit': self.ml.pk, 'is_active': 'on'})
        self.assertEqual(r.status_code, 302)
        # but nothing already entered can be changed
        self.adder.post(reverse('link_update', args=[self.link.pk]), {
            f'l{self.link.pk}-quantity': '99', f'l{self.link.pk}-unit': self.ml.pk, f'l{self.link.pk}-is_required': 'on',
            f'l{self.link.pk}-notes': ''})
        self.link.refresh_from_db()
        self.assertEqual(self.link.quantity, Decimal('3'))
        self.assertEqual(self.adder.get(reverse('consumable_edit', args=[self.contrast.pk])).status_code, 403)
        self.adder.post(reverse('consumable_toggle', args=[self.contrast.pk]))
        self.contrast.refresh_from_db()
        self.assertTrue(self.contrast.is_active)
        self.assertEqual(self.adder.get(reverse('unit_list')).status_code, 200)

    def test_adder_sees_values_as_text_and_no_edit_buttons(self):
        html = self.adder.get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        self.assertNotIn('name="l%d-quantity"' % self.link.pk, html)
        self.assertNotIn('>Save<', html)
        self.assertIn('Add consumables', html)
        self.assertNotIn('Edit</a>', self.adder.get(reverse('consumable_list')).content.decode())

    def test_adder_cannot_replace_in_bulk(self):
        data = {'exams': [self.exam.pk], 'row_consumable': [self.contrast.pk], 'row_quantity': ['7'],
                'row_unit': [self.ml.pk], 'row_required': ['1'], 'row_notes': [''], 'conflict': 'overwrite'}
        self.assertEqual(self.adder.post(reverse('bulk_apply'), data).status_code, 200)
        self.link.refresh_from_db()
        self.assertEqual(self.link.quantity, Decimal('3'))
        data['conflict'] = 'skip'
        self.assertEqual(self.adder.post(reverse('bulk_apply'), data).status_code, 302)

    def test_admin_can_grant_and_revoke(self):
        u = make_user('later', edit=True)
        c = client_for(u)
        self.assertEqual(c.get(reverse('consumable_edit', args=[self.contrast.pk])).status_code, 403)
        self.client.post(reverse('user_edit', args=[u.pk]), {
            'full_name': '', 'account_type': 'user', 'can_edit': 'on', 'can_change': 'on', 'is_active': 'on'})
        self.assertEqual(c.get(reverse('consumable_edit', args=[self.contrast.pk])).status_code, 200)

    def test_change_needs_add(self):
        r = self.client.post(reverse('user_edit', args=[make_user('z').pk]), {
            'full_name': '', 'account_type': 'user', 'can_change': 'on', 'is_active': 'on'})
        self.assertEqual(r.status_code, 200)

    def test_admin_edits_entered_consumables(self):
        self.client.post(reverse('link_update', args=[self.link.pk]), {
            f'l{self.link.pk}-quantity': '5', f'l{self.link.pk}-unit': self.ml.pk, f'l{self.link.pk}-is_required': '',
            f'l{self.link.pk}-notes': 'checked'})
        self.link.refresh_from_db()
        self.assertEqual((self.link.quantity, self.link.is_required, self.link.notes), (Decimal('5'), False, 'checked'))
        self.client.post(reverse('consumable_edit', args=[self.contrast.pk]), {
            'name': 'Contrast 350', 'default_unit': self.ml.pk, 'is_active': 'on'})
        self.contrast.refresh_from_db()
        self.assertEqual(self.contrast.name, 'Contrast 350')


class Alternatives(BaseCase):
    def add(self, name='Omni', alts=(), **extra):
        data = {'name': name, 'default_unit': self.ml.pk, 'is_active': 'on', 'alt_name': list(alts) or [''], **extra}
        return self.client.post(reverse('consumable_create'), data)

    def test_create_with_alternatives_creates_missing_ones(self):
        self.add('Omni', ['Visi 320', 'contrast', '  Iso   Vue ', 'Omni', ''])      # 'contrast' exists (any capitals)
        omni = Consumable.objects.get(name='Omni')
        names = sorted(l.alternative.name for l in omni.alternative_links.all())
        self.assertEqual(names, ['Contrast', 'Iso Vue', 'Visi 320'])
        self.assertEqual(Consumable.objects.filter(name__iexact='contrast').count(), 1)
        new = Consumable.objects.get(name='Visi 320')
        self.assertEqual(new.default_unit, self.ml)                                  # took the primary's unit

    def test_typed_names_use_the_excel_list_details_when_present(self):
        cat = ConsumableCategory.objects.get_or_create(name='Contrast media')[0]
        CatalogItem.objects.create(name='Visi 320', category=cat, default_unit=self.piece, description='from list')
        self.add('Omni', ['visi 320'])
        alt = Consumable.objects.get(name='visi 320')
        self.assertEqual((alt.category, alt.default_unit, alt.description), (cat, self.piece, 'from list'))

    def test_edit_replaces_the_alternatives(self):
        self.add('Omni', ['A1', 'A2'])
        omni = Consumable.objects.get(name='Omni')
        self.client.post(reverse('consumable_edit', args=[omni.pk]), {
            'name': 'Omni', 'default_unit': self.ml.pk, 'is_active': 'on', 'alt_name': ['A2', 'A3']})
        self.assertEqual(sorted(l.alternative.name for l in omni.alternative_links.all()), ['A2', 'A3'])
        self.assertTrue(Consumable.objects.filter(name='A1').exists())               # only the link was removed
        page = self.client.get(reverse('consumable_edit', args=[omni.pk])).content.decode()
        self.assertIn('A2', page)

    def test_alternatives_shown_on_examination_and_list(self):
        self.add('Omni', ['A1'])
        omni = Consumable.objects.get(name='Omni')
        ExaminationConsumable.objects.create(examination=self.exam, consumable=omni, quantity=1, unit=self.ml)
        self.assertContains(self.client.get(reverse('exam_detail', args=[self.exam.pk])), 'Alternatives:')
        self.assertContains(self.client.get(reverse('consumable_list')), 'Alternatives: A1')
        viewer = client_for(make_user('v2'))
        page = viewer.get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        self.assertIn('A1', page)
        self.assertNotIn('instead</option>', page)

    def test_swap_to_an_alternative(self):
        self.add('Omni', ['A1'])
        omni, a1 = Consumable.objects.get(name='Omni'), Consumable.objects.get(name='A1')
        link = ExaminationConsumable.objects.create(examination=self.exam, consumable=omni, quantity=4, unit=self.ml,
                                                    notes='n')
        self.assertContains(self.client.get(reverse('exam_detail', args=[self.exam.pk])), 'Use A1 instead')
        self.client.post(reverse('link_swap', args=[link.pk]), {'alternative': self.cannula.pk})   # not an alternative
        link.refresh_from_db()
        self.assertEqual(link.consumable, omni)
        self.client.post(reverse('link_swap', args=[link.pk]), {'alternative': a1.pk})
        link.refresh_from_db()
        self.assertEqual((link.consumable, link.quantity, link.notes), (a1, Decimal('4'), 'n'))
        # refused when the alternative is already on the examination, and for people who may not change entries
        other = ExaminationConsumable.objects.create(examination=self.exam, consumable=omni, quantity=1, unit=self.ml)
        self.client.post(reverse('link_swap', args=[other.pk]), {'alternative': a1.pk})
        other.refresh_from_db()
        self.assertEqual(other.consumable, omni)
        adder = client_for(make_user('adder2', edit=True))
        adder.post(reverse('link_swap', args=[other.pk]), {'alternative': a1.pk})
        other.refresh_from_db()
        self.assertEqual(other.consumable, omni)

    def test_alternative_in_use_is_retired_not_deleted(self):
        self.add('Omni', ['A1'])
        a1 = Consumable.objects.get(name='A1')
        self.assertTrue(a1.is_in_use)
        self.client.post(reverse('consumable_delete', args=[a1.pk]))
        a1.refresh_from_db()
        self.assertFalse(a1.is_active)

    def test_deleting_a_primary_removes_only_its_links(self):
        self.add('Omni', ['A1'])
        omni = Consumable.objects.get(name='Omni')
        self.client.post(reverse('consumable_delete', args=[omni.pk]))
        self.assertFalse(Consumable.objects.filter(name='Omni').exists())
        self.assertTrue(Consumable.objects.filter(name='A1').exists())

    def test_excel_export_has_alternatives_column(self):
        from io import BytesIO
        from openpyxl import load_workbook
        self.add('Omni', ['A1', 'A2'])
        omni = Consumable.objects.get(name='Omni')
        ExaminationConsumable.objects.create(examination=self.exam, consumable=omni, quantity=1, unit=self.ml)
        wb = load_workbook(BytesIO(self.client.get(reverse('export_excel'), {'download': '1'}).content))
        header = [c.value for c in wb['Examination consumables'][1]]
        row = [r for r in wb['Examination consumables'].iter_rows(min_row=2, values_only=True) if r[5] == 'Omni'][0]
        self.assertEqual(row[header.index('Alternatives')], 'A1; A2')

    def test_suggest_for_alternatives_offers_existing_consumables_first(self):
        CatalogItem.objects.create(name='Contrast agent', default_unit=self.ml)
        r = self.client.get(reverse('consumable_suggest'), {'q': 'contr', 'kind': 'alt'}).json()
        self.assertEqual([(i['name'], i['source']) for i in r['results']],
                         [('Contrast', 'consumable'), ('Contrast agent', 'list')])
        self.assertEqual(r['total'], Consumable.objects.count() + 1)
        plain = self.client.get(reverse('consumable_suggest'), {'q': 'contr'}).json()
        self.assertEqual([i['name'] for i in plain['results']], ['Contrast agent'])

    def test_total_zero_without_a_list_so_the_form_does_not_cry_red(self):
        self.assertEqual(self.client.get(reverse('consumable_suggest'), {'q': 'anything'}).json(),
                         {'results': [], 'total': 0})

    def test_form_markup(self):
        page = self.client.get(reverse('consumable_create')).content.decode()
        for needle in ('id="alt-rows"', 'id="add-alt"', 'suggest.js', 'data-suggest-url'):
            self.assertIn(needle, page)


class ListAlternativesColumn(BaseCase):
    def test_import_reads_alternatives_and_add_all_links_them(self):
        from django.core.files.uploadedfile import SimpleUploadedFile
        data = excel_bytes({'S': [['Name', 'Unit', 'Alternatives'],
                                  ['Gel A', 'ml', 'Gel B; Gel C'], ['Gel B', 'ml', '']]})
        self.client.post(reverse('catalog_import'), {'file': SimpleUploadedFile('l.xlsx', data), 'add_all': '1'})
        item = CatalogItem.objects.get(name='Gel A')
        self.assertEqual(item.alternatives, 'Gel B; Gel C')
        gel_a = Consumable.objects.get(name='Gel A')
        self.assertEqual(sorted(l.alternative.name for l in gel_a.alternative_links.all()), ['Gel B', 'Gel C'])
        sugg = self.client.get(reverse('consumable_suggest'), {'q': 'gel a'}).json()['results']
        self.assertEqual(sugg[0]['alternatives'], ['Gel B', 'Gel C'])


class SavedSets(BaseCase):
    def setUp(self):
        super().setUp()
        self.glove = Consumable.objects.create(name='Gloves', default_unit=self.piece)

    def make_set(self, name='Standard CT'):
        s = ConsumableSet.objects.create(name=name)
        ConsumableSetItem.objects.create(consumable_set=s, consumable=self.contrast, quantity=100, unit=self.ml)
        ConsumableSetItem.objects.create(consumable_set=s, consumable=self.glove, quantity=2, unit=self.piece,
                                         is_required=False)
        return s

    def test_build_a_set_independently(self):
        r = self.client.post(reverse('set_create'), {'name': 'IV start', 'description': 'basic'})
        s = ConsumableSet.objects.get(name='IV start')
        self.assertEqual(r.status_code, 302)
        self.client.post(reverse('set_item_add', args=[s.pk]), {
            'consumable': self.cannula.pk, 'quantity': '1', 'unit': self.piece.pk, 'is_required': 'on'})
        self.client.post(reverse('set_item_add', args=[s.pk]), {
            'consumable': self.cannula.pk, 'quantity': '1', 'unit': self.piece.pk})          # duplicate
        self.client.post(reverse('set_item_add', args=[s.pk]), {
            'consumable': self.glove.pk, 'quantity': '1.5', 'unit': self.piece.pk})          # not whole
        self.assertEqual(s.items.count(), 1)
        self.assertEqual(self.client.post(reverse('set_create'), {'name': 'iv START'}).status_code, 200)   # unique
        self.assertEqual(ConsumableSet.objects.count(), 1)

    def test_apply_set_to_examination(self):
        s = self.make_set()
        self.client.post(reverse('exam_apply_set', args=[self.exam.pk]), {'set': s.pk})
        self.assertEqual(self.exam.consumable_links.count(), 2)
        got = self.exam.consumable_links.get(consumable=self.glove)
        self.assertEqual((got.quantity, got.is_required), (Decimal('2'), False))
        # applying again changes nothing; an inactive consumable is left out
        self.exam.consumable_links.filter(consumable=self.glove).update(quantity=9)
        self.glove.is_active = False
        self.glove.save()
        self.client.post(reverse('exam_apply_set', args=[self.exam2.pk]), {'set': s.pk})
        self.assertEqual(list(self.exam2.consumable_links.values_list('consumable__name', flat=True)), ['Contrast'])
        self.client.post(reverse('exam_apply_set', args=[self.exam.pk]), {'set': s.pk})
        self.assertEqual(self.exam.consumable_links.get(consumable=self.glove).quantity, Decimal('9'))

    def test_save_examination_as_set_and_reuse(self):
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast, quantity=3, unit=self.ml)
        self.client.post(reverse('exam_save_as_set', args=[self.exam.pk]), {'name': 'From CT'})
        s = ConsumableSet.objects.get(name='From CT')
        self.assertEqual(s.items.get().quantity, Decimal('3'))
        self.client.post(reverse('exam_save_as_set', args=[self.exam.pk]), {'name': 'from ct'})        # same name
        self.assertEqual(ConsumableSet.objects.count(), 1)
        self.client.post(reverse('exam_save_as_set', args=[self.exam2.pk]), {'name': 'Empty one'})     # nothing to save
        self.assertEqual(ConsumableSet.objects.count(), 1)
        self.client.post(reverse('exam_apply_set', args=[self.exam2.pk]), {'set': s.pk})
        self.assertEqual(self.exam2.consumable_links.get().quantity, Decimal('3'))

    def test_bulk_apply_with_a_set(self):
        s = self.make_set()
        data = {'exams': [self.exam.pk, self.exam2.pk], 'set': s.pk, 'conflict': 'skip',
                'row_consumable': [self.cannula.pk, ''], 'row_quantity': ['4', '1'], 'row_unit': [self.piece.pk, ''],
                'row_required': ['1', '1'], 'row_notes': ['', '']}
        self.assertEqual(self.client.post(reverse('bulk_apply'), data).status_code, 302)
        for e in (self.exam, self.exam2):
            self.assertEqual(e.consumable_links.count(), 3)
        self.assertEqual(self.client.get(reverse('bulk_apply'), {'set': s.pk}).context['posted_set'], str(s.pk))
        # a set alone is enough
        data2 = {'exams': [self.exam.pk], 'set': s.pk, 'row_consumable': [''], 'row_quantity': ['1'],
                 'row_unit': [''], 'row_required': ['1'], 'row_notes': ['']}
        self.assertEqual(self.client.post(reverse('bulk_apply'), data2).status_code, 302)

    def test_permissions(self):
        s = self.make_set()
        viewer = client_for(make_user('sv'))
        self.assertEqual(viewer.get(reverse('set_list')).status_code, 200)
        page = viewer.get(reverse('set_detail', args=[s.pk])).content.decode()
        self.assertIn('Contrast', page)
        self.assertNotIn('Add consumables to the set', page)
        self.assertEqual(viewer.get(reverse('set_create')).status_code, 403)
        adder = client_for(make_user('sa', edit=True))
        self.assertEqual(adder.post(reverse('set_create'), {'name': 'Mine'}).status_code, 302)
        item = s.items.first()
        adder.post(reverse('set_item_update', args=[item.pk]), {
            f's{item.pk}-quantity': '55', f's{item.pk}-unit': self.ml.pk, f's{item.pk}-is_required': 'on', f's{item.pk}-notes': ''})
        item.refresh_from_db()
        self.assertEqual(item.quantity, Decimal('100'))                       # changing needs the change permission
        adder.post(reverse('set_toggle', args=[s.pk]))
        adder.post(reverse('set_item_remove', args=[item.pk]))
        adder.post(reverse('set_delete', args=[s.pk]))
        s.refresh_from_db()
        self.assertTrue(s.is_active)
        self.assertEqual(s.items.count(), 2)
        changer = client_for(make_user('sc', edit=True, change=True))
        changer.post(reverse('set_item_update', args=[item.pk]), {
            f's{item.pk}-quantity': '55', f's{item.pk}-unit': self.ml.pk, f's{item.pk}-is_required': 'on', f's{item.pk}-notes': ''})
        item.refresh_from_db()
        self.assertEqual(item.quantity, Decimal('55'))
        changer.post(reverse('set_item_remove', args=[item.pk]))               # still needs delete
        self.assertEqual(s.items.count(), 2)
        self.client.post(reverse('set_item_remove', args=[item.pk]))
        self.assertEqual(s.items.count(), 1)

    def test_deleting_a_set_keeps_what_was_applied(self):
        s = self.make_set()
        self.client.post(reverse('exam_apply_set', args=[self.exam.pk]), {'set': s.pk})
        self.client.post(reverse('set_delete', args=[s.pk]))
        self.assertFalse(ConsumableSet.objects.exists())
        self.assertEqual(self.exam.consumable_links.count(), 2)

    def test_consumable_in_a_set_is_retired_not_deleted(self):
        self.make_set()
        self.client.post(reverse('consumable_delete', args=[self.glove.pk]))
        self.glove.refresh_from_db()
        self.assertFalse(self.glove.is_active)

    def test_pages_and_export_sheet(self):
        from io import BytesIO
        from openpyxl import load_workbook
        s = self.make_set()
        for url in (reverse('set_list'), reverse('set_detail', args=[s.pk]), reverse('set_create'),
                    reverse('set_edit', args=[s.pk]), reverse('exam_detail', args=[self.exam.pk])):
            self.assertEqual(self.client.get(url).status_code, 200, url)
        self.assertContains(self.client.get(reverse('exam_detail', args=[self.exam.pk])), 'Apply a saved set')
        wb = load_workbook(BytesIO(self.client.get(reverse('export_excel'), {'download': '1'}).content))
        rows = list(wb['Saved sets'].iter_rows(min_row=2, values_only=True))
        self.assertIn(('Standard CT', 'Contrast', 100, 'ml', 'Required', None), rows)

    def test_whole_number_rule_applies_to_sets(self):
        s = ConsumableSet.objects.create(name='X')
        item = ConsumableSetItem(consumable_set=s, consumable=self.contrast, quantity=Decimal('1.5'), unit=self.ml)
        from django.core.exceptions import ValidationError
        with self.assertRaises(ValidationError):
            item.full_clean()


class RemoveConsumableEverywhere(BaseCase):
    def setUp(self):
        super().setUp()
        us = Modality.objects.get(slug='ultrasound')
        self.us_exams = [Examination.objects.create(name=f'US {i}', modality=us) for i in range(3)]
        for e in self.us_exams + [self.exam]:
            ExaminationConsumable.objects.create(examination=e, consumable=self.contrast, quantity=1, unit=self.ml)
        ExaminationConsumable.objects.create(examination=self.us_exams[0], consumable=self.cannula, quantity=1, unit=self.piece)

    def test_usage_page_counts_by_modality(self):
        r = self.client.get(reverse('consumable_usage', args=[self.contrast.pk]))
        self.assertEqual(r.context['total'], 4)
        self.assertEqual({x['slug']: x['n'] for x in r.context['rows']}, {'ct': 1, 'ultrasound': 3})
        self.assertContains(r, 'Ultrasound only (3)')

    def test_remove_from_one_modality_only(self):
        self.client.post(reverse('consumable_remove_everywhere', args=[self.contrast.pk]), {'modality': 'ultrasound'})
        self.assertEqual(ExaminationConsumable.objects.filter(consumable=self.contrast).count(), 1)
        self.assertTrue(self.exam.consumable_links.filter(consumable=self.contrast).exists())
        self.assertEqual(ExaminationConsumable.objects.filter(consumable=self.cannula).count(), 1)   # others untouched

    def test_remove_everywhere_then_it_can_be_deleted(self):
        self.client.post(reverse('consumable_remove_everywhere', args=[self.contrast.pk]), {'modality': ''})
        self.assertFalse(ExaminationConsumable.objects.filter(consumable=self.contrast).exists())
        self.assertFalse(Consumable.objects.get(pk=self.contrast.pk).is_in_use)
        self.client.post(reverse('consumable_delete', args=[self.contrast.pk]))        # now a real delete
        self.assertFalse(Consumable.objects.filter(pk=self.contrast.pk).exists())

    def test_remove_and_delete_in_one_step(self):
        self.client.post(reverse('consumable_remove_everywhere', args=[self.contrast.pk]), {'also_delete': '1'})
        self.assertFalse(Consumable.objects.filter(pk=self.contrast.pk).exists())

    def test_also_delete_keeps_it_when_a_set_uses_it(self):
        s = ConsumableSet.objects.create(name='S')
        ConsumableSetItem.objects.create(consumable_set=s, consumable=self.contrast, quantity=1, unit=self.ml)
        self.client.post(reverse('consumable_remove_everywhere', args=[self.contrast.pk]), {'also_delete': '1'})
        self.assertTrue(Consumable.objects.filter(pk=self.contrast.pk).exists())
        self.assertFalse(ExaminationConsumable.objects.filter(consumable=self.contrast).exists())

    def test_needs_delete_permission(self):
        for u in (make_user('a1', edit=True), make_user('a2', edit=True, change=True)):
            c = client_for(u)
            self.assertEqual(c.get(reverse('consumable_usage', args=[self.contrast.pk])).status_code, 403)
            c.post(reverse('consumable_remove_everywhere', args=[self.contrast.pk]), {})
        self.assertEqual(ExaminationConsumable.objects.filter(consumable=self.contrast).count(), 4)
        deleter = client_for(make_user('a3', edit=True, delete=True))
        self.assertEqual(deleter.get(reverse('consumable_usage', args=[self.contrast.pk])).status_code, 200)
        self.assertContains(deleter.get(reverse('consumable_list')), 'Remove from exams')
        deleter.post(reverse('consumable_remove_everywhere', args=[self.contrast.pk]), {})
        self.assertEqual(ExaminationConsumable.objects.filter(consumable=self.contrast).count(), 0)


class ExamViewExport(BaseCase):
    """Excel 'exam view': examination (+ machine code) merged on the left, required consumables in rows beside it,
    optional items as replacements."""

    def setUp(self):
        super().setUp()
        self.exam.code = 'L00761'
        self.exam.save()
        self.pair = Unit.objects.get(name='pair')
        self.gloves = Consumable.objects.create(name='Gloves', default_unit=self.pair)
        self.gloves_pf = Consumable.objects.create(name='Gloves powder free', default_unit=self.pair)
        self.mask = Consumable.objects.create(name='Surgical mask', default_unit=self.piece)
        self.syr = Consumable.objects.create(name='Syringe 20cc', default_unit=self.piece)
        self.syr20 = Consumable.objects.create(name='Syringe 20 ml', default_unit=self.piece)
        self.contrast2 = Consumable.objects.create(name='Contrast B', default_unit=self.ml)

    def link(self, c, qty, unit, required=True, replaces=None, notes='', exam=None):
        return ExaminationConsumable.objects.create(examination=exam or self.exam, consumable=c, quantity=qty,
                                                    unit=unit, is_required=required, replaces=replaces, notes=notes)

    def workbook(self, **params):
        from io import BytesIO
        from openpyxl import load_workbook
        r = self.client.get(reverse('export_excel'), {'download': '1', 'layout': 'exam', **params})
        self.assertEqual(r.status_code, 200)
        return r, load_workbook(BytesIO(r.content)).active

    def test_layout_merged_cells_code_and_replacements(self):
        self.link(self.contrast, 90, self.ml)
        self.link(self.contrast2, 50, self.ml, required=False, replaces=self.contrast)          # set on the exam
        self.link(self.gloves, 3, self.pair)
        self.link(self.gloves_pf, 3, self.pair, required=False, notes='for allergic staff')    # via alternatives
        ConsumableAlternative.objects.create(consumable=self.gloves, alternative=self.gloves_pf)
        self.link(self.cannula, 2, self.piece)
        r, ws = self.workbook(modality='ct')
        self.assertIn('exam-view', r['Content-Disposition'])
        header = [c.value for c in ws[1]]
        self.assertEqual(header[:7], ['Modality', 'Examination', 'Machine code', 'Required consumable', 'Qty', 'Unit',
                                      'Replacement if required item is unavailable'])
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual(len(rows), 3)                                      # only the 3 required items get rows
        self.assertEqual(rows[0][:3], ('CT Scan', 'CT Abdomen/Pelvis with contrast', 'L00761'))
        self.assertEqual(rows[1][:3], (None, None, None))                   # merged: one cell per examination
        self.assertEqual([str(m) for m in ws.merged_cells.ranges if m.min_col == 2], ['B2:B4'])
        self.assertEqual([str(m) for m in ws.merged_cells.ranges if m.min_col == 3], ['C2:C4'])
        by_name = {r[3]: r for r in rows}
        self.assertEqual(by_name['Contrast'][4:7], (90, 'ml', 'Contrast B (50 ml)'))
        self.assertEqual(by_name['Gloves'][6], 'Gloves powder free (3 pair) - for allergic staff')
        self.assertEqual(by_name['IV Cannula'][6], '-')
        self.assertNotIn('Other optional items', ' '.join(str(h) for h in header))

    def test_required_rows_are_alphabetical_and_options_sorted(self):
        self.link(self.cannula, 1, self.piece)
        self.link(self.contrast, 1, self.ml)
        self.link(self.gloves_pf, 1, self.pair, required=False, replaces=self.contrast)
        self.link(self.contrast2, 1, self.ml, required=False, replaces=self.contrast)
        _, ws = self.workbook(modality='ct')
        rows = list(ws.iter_rows(min_row=2, values_only=True))
        self.assertEqual([r[3] for r in rows], ['Contrast', 'IV Cannula'])
        self.assertEqual(rows[0][6], 'Contrast B (1 ml)\nGloves powder free (1 pair)')

    def test_unlinked_optional_items_are_not_lost(self):
        self.link(self.gloves, 3, self.pair)
        self.link(self.mask, 1, self.piece, required=False)
        _, ws = self.workbook(modality='ct')
        header = [c.value for c in ws[1]]
        self.assertEqual(header[7], 'Other optional items (not linked to a required item)')
        self.assertEqual(ws.cell(2, 8).value, 'Surgical mask (1 piece)')

    def test_explicit_link_beats_alternatives_and_stale_links_fall_back(self):
        self.link(self.syr, 4, self.piece)
        self.link(self.contrast, 1, self.ml)
        ConsumableAlternative.objects.create(consumable=self.syr, alternative=self.syr20)
        self.link(self.syr20, 5, self.piece, required=False, replaces=self.contrast)    # explicit: replaces contrast
        _, ws = self.workbook(modality='ct')
        by = {r[3]: r for r in ws.iter_rows(min_row=2, values_only=True)}
        self.assertEqual(by['Contrast'][6], 'Syringe 20 ml (5 piece)')
        self.assertEqual(by['Syringe 20cc'][6], '-')

    def test_exam_with_only_optional_items_and_empty_exams(self):
        self.link(self.mask, 1, self.piece, required=False)
        _, ws = self.workbook(modality='ct')
        row = list(ws.iter_rows(min_row=2, values_only=True))[0]
        self.assertEqual((row[2], row[3], row[7]), ('L00761', None, 'Surgical mask (1 piece)'))
        _, with_empty = self.workbook(modality='mri')                       # MRI Brain has no consumables
        self.assertEqual(with_empty.cell(2, 4).value, '(no consumables entered yet)')
        _, without = self.workbook(modality='mri', empty='0')
        self.assertEqual(without.max_row, 1)

    def test_inactive_and_modality_filters_and_ordering(self):
        self.link(self.gloves, 1, self.pair)
        self.exam2.is_active = False
        self.exam2.save()
        self.link(self.gloves, 1, self.pair, exam=self.exam2)
        _, ws = self.workbook()
        names = [r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r[1]]
        self.assertEqual(names, ['CT Abdomen/Pelvis with contrast'])
        _, ws = self.workbook(inactive='1')
        names = [r[1] for r in ws.iter_rows(min_row=2, values_only=True) if r[1]]
        self.assertEqual(names, ['CT Abdomen/Pelvis with contrast', 'MRI Brain'])    # CT before MRI

    def test_text_that_looks_like_a_formula_stays_text(self):
        odd = Consumable.objects.create(name='=SUM(1,1)', default_unit=self.ml)
        self.link(odd, 1, self.ml)
        _, ws = self.workbook(modality='ct')
        cell = [c for c in ws['D'] if c.value == '=SUM(1,1)'][0]
        self.assertEqual(cell.data_type, 's')

    def test_default_download_is_still_the_full_workbook_and_page_offers_both(self):
        from io import BytesIO
        from openpyxl import load_workbook
        r = self.client.get(reverse('export_excel'), {'download': '1'})
        self.assertEqual(load_workbook(BytesIO(r.content)).sheetnames[0], 'Modalities')
        page = self.client.get(reverse('export_excel')).content.decode()
        self.assertIn('value="exam" checked', page)
        self.assertIn('value="full"', page)

    def test_viewers_can_download_it(self):
        self.link(self.gloves, 1, self.pair)
        r = client_for(make_user('exp')).get(reverse('export_excel'), {'download': '1', 'layout': 'exam'})
        self.assertEqual(r.status_code, 200)


class ReplacesField(BaseCase):
    def setUp(self):
        super().setUp()
        self.glove = Consumable.objects.create(name='Gloves', default_unit=self.piece)
        self.alt = Consumable.objects.create(name='Gloves M', default_unit=self.piece)
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.glove, quantity=2, unit=self.piece)

    def add_optional(self, **extra):
        return self.client.post(reverse('link_add', args=[self.exam.pk]), {
            'consumable': self.alt.pk, 'quantity': '2', 'unit': self.piece.pk, 'replaces': self.glove.pk, **extra})

    def test_optional_item_can_be_set_to_replace_a_required_one(self):
        self.add_optional()
        link = self.exam.consumable_links.get(consumable=self.alt)
        self.assertEqual((link.is_required, link.replaces), (False, self.glove))
        page = client_for(make_user('view')).get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        self.assertIn('replaces Gloves', page)

    def test_required_items_never_keep_a_replaces_and_it_cannot_point_at_an_optional_item(self):
        self.add_optional(is_required='on')
        self.assertIsNone(self.exam.consumable_links.get(consumable=self.alt).replaces)   # required: link dropped
        opt = Consumable.objects.create(name='Gloves S', default_unit=self.piece)
        self.client.post(reverse('link_add', args=[self.exam.pk]), {
            'consumable': opt.pk, 'quantity': '1', 'unit': self.piece.pk, 'replaces': self.glove.pk})
        self.assertEqual(self.exam.consumable_links.get(consumable=opt).replaces, self.glove)
        other = Consumable.objects.create(name='Mask', default_unit=self.piece)
        self.client.post(reverse('link_add', args=[self.exam.pk]), {
            'consumable': other.pk, 'quantity': '1', 'unit': self.piece.pk, 'replaces': opt.pk})
        self.assertFalse(self.exam.consumable_links.filter(consumable=other).exists())    # opt is optional: not allowed

    def test_update_changes_the_link_and_swapping_or_removing_keeps_things_tidy(self):
        self.add_optional()
        link = self.exam.consumable_links.get(consumable=self.alt)
        self.client.post(reverse('link_update', args=[link.pk]), {
            f'l{link.pk}-quantity': '3', f'l{link.pk}-unit': self.piece.pk, f'l{link.pk}-replaces': '',
            f'l{link.pk}-notes': ''})
        link.refresh_from_db()
        self.assertIsNone(link.replaces)
        self.client.post(reverse('link_update', args=[link.pk]), {
            f'l{link.pk}-quantity': '3', f'l{link.pk}-unit': self.piece.pk, f'l{link.pk}-replaces': self.glove.pk,
            f'l{link.pk}-notes': ''})
        link.refresh_from_db()
        self.assertEqual(link.replaces, self.glove)
        req = self.exam.consumable_links.get(consumable=self.glove)
        self.client.post(reverse('link_remove', args=[req.pk]))                  # the required item goes
        link.refresh_from_db()
        self.assertIsNone(link.replaces)

    def test_swap_moves_replacements_to_the_new_item(self):
        newer = Consumable.objects.create(name='Gloves XL', default_unit=self.piece)
        ConsumableAlternative.objects.create(consumable=self.glove, alternative=newer)
        self.add_optional()
        req = self.exam.consumable_links.get(consumable=self.glove)
        self.client.post(reverse('link_swap', args=[req.pk]), {'alternative': newer.pk})
        self.assertEqual(self.exam.consumable_links.get(consumable=self.alt).replaces, newer)

    def test_copy_and_bulk_apply_carry_the_link(self):
        self.add_optional()
        self.client.post(reverse('exam_copy', args=[self.exam2.pk]), {'source': self.exam.pk})
        copied = self.exam2.consumable_links.get(consumable=self.alt)
        self.assertEqual((copied.is_required, copied.replaces), (False, self.glove))


class SaveEverythingWithOneClick(BaseCase):
    """One Save changes button per page: all edited rows plus a new item, saved together or not at all."""

    def setUp(self):
        super().setUp()
        self.glove = Consumable.objects.create(name='Gloves', default_unit=self.piece)
        self.mask = Consumable.objects.create(name='Mask', default_unit=self.piece)
        self.l1 = ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast, quantity=3, unit=self.ml)
        self.l2 = ExaminationConsumable.objects.create(examination=self.exam, consumable=self.cannula, quantity=1, unit=self.piece)
        self.url = reverse('exam_save_all', args=[self.exam.pk])

    def row(self, link, qty, unit=None, required='on', notes='', replaces=''):
        p = f'l{link.pk}-'
        return {p + 'quantity': qty, p + 'unit': (unit or link.unit).pk, p + 'is_required': required,
                p + 'replaces': replaces, p + 'notes': notes}

    def test_one_click_saves_every_edited_row_and_a_new_item(self):
        data = {**self.row(self.l1, '9', notes='checked'), **self.row(self.l2, '4'),
                'new0-consumable': self.glove.pk, 'new0-quantity': '2', 'new0-unit': self.piece.pk, 'new0-is_required': 'on', 'new0-notes': 'new'}
        r = self.client.post(self.url, data)
        self.assertEqual(r.status_code, 302)
        self.l1.refresh_from_db(); self.l2.refresh_from_db()
        self.assertEqual((self.l1.quantity, self.l1.notes, self.l2.quantity), (Decimal('9'), 'checked', Decimal('4')))
        added = self.exam.consumable_links.get(consumable=self.glove)
        self.assertEqual((added.quantity, added.notes), (Decimal('2'), 'new'))
        msg = [str(m) for m in self.client.get(reverse('exam_detail', args=[self.exam.pk])).context['messages']]
        self.assertIn('Saved: 2 rows updated, 1 added (Gloves).', msg)

    def test_nothing_is_saved_when_any_part_is_wrong_and_the_typed_values_come_back(self):
        data = {**self.row(self.l1, '9'), **self.row(self.l2, '1.5'),                  # 1.5 piece: not allowed
                'new0-consumable': self.glove.pk, 'new0-quantity': '2', 'new0-unit': self.piece.pk, 'new0-is_required': 'on'}
        r = self.client.post(self.url, data)
        self.assertEqual(r.status_code, 200)
        self.l1.refresh_from_db()
        self.assertEqual(self.l1.quantity, Decimal('3'))                              # the good row was not saved either
        self.assertFalse(self.exam.consumable_links.filter(consumable=self.glove).exists())
        html = r.content.decode()
        self.assertIn('value="9"', html)                                               # what was typed is shown again
        self.assertIn('value="1.5"', html)
        self.assertIn('Use a whole number', html)
        self.assertIn('data-dirty-start="1"', html)
        self.assertIn('Nothing was saved', ' '.join(str(m) for m in r.context['messages']))

    def test_a_bad_new_item_also_blocks_the_row_changes(self):
        data = {**self.row(self.l1, '9'), 'new0-consumable': self.contrast.pk, 'new0-quantity': '1', 'new0-unit': self.ml.pk,
                'new0-is_required': 'on'}                                                   # already on the examination
        self.assertEqual(self.client.post(self.url, data).status_code, 200)
        self.l1.refresh_from_db()
        self.assertEqual(self.l1.quantity, Decimal('3'))

    def test_pressing_save_without_changes_says_so(self):
        r = self.client.post(self.url, {**self.row(self.l1, '3'), **self.row(self.l2, '1')}, follow=True)
        self.assertIn('No changes to save.', [str(m) for m in r.context['messages']])

    def test_new_item_only(self):
        self.client.post(self.url, {'new0-consumable': self.mask.pk, 'new0-quantity': '5', 'new0-unit': self.piece.pk,
                                    'new0-is_required': '', 'new0-replaces': self.contrast.pk})
        link = self.exam.consumable_links.get(consumable=self.mask)
        self.assertEqual((link.is_required, link.replaces, link.quantity), (False, self.contrast, Decimal('5')))

    def test_required_optional_and_replaces_are_saved_together(self):
        self.client.post(self.url, {**self.row(self.l1, '3'), **self.row(self.l2, '1', required='', replaces=self.contrast.pk)})
        self.l2.refresh_from_db()
        self.assertEqual((self.l2.is_required, self.l2.replaces), (False, self.contrast))

    def test_a_user_who_may_only_add_cannot_sneak_in_row_changes(self):
        adder = client_for(make_user('adder9', edit=True))
        r = adder.post(self.url, {**self.row(self.l1, '99'), 'new0-consumable': self.glove.pk, 'new0-quantity': '1',
                                  'new0-unit': self.piece.pk, 'new0-is_required': 'on'})
        self.assertEqual(r.status_code, 302)
        self.l1.refresh_from_db()
        self.assertEqual(self.l1.quantity, Decimal('3'))
        self.assertFalse(self.exam.consumable_links.filter(consumable=self.glove).exists())        # nothing at all saved
        adder.post(self.url, {'new0-consumable': self.glove.pk, 'new0-quantity': '1', 'new0-unit': self.piece.pk, 'new0-is_required': 'on'})
        self.assertTrue(self.exam.consumable_links.filter(consumable=self.glove).exists())         # adding alone is fine

    def test_viewers_cannot_use_it(self):
        v = client_for(make_user('viewer9'))
        v.post(self.url, {**self.row(self.l1, '99')})
        self.l1.refresh_from_db()
        self.assertEqual(self.l1.quantity, Decimal('3'))

    def test_page_has_one_form_one_button_and_no_per_row_save(self):
        html = self.client.get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        self.assertEqual(html.count('id="examform"'), 1)
        self.assertIn('id="savebar"', html)
        self.assertNotIn('/exam-consumables/%d/update/' % self.l1.pk, html)                 # no separate row-save forms
        self.assertIn('name="l%d-quantity"' % self.l1.pk, html)
        self.assertIn('form="examform"', html)
        adder = client_for(make_user('adder10', edit=True)).get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        self.assertIn('id="savebar"', adder)
        viewer = client_for(make_user('viewer10')).get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        self.assertNotIn('id="savebar"', viewer)
        self.assertNotIn('examform', viewer)

    def test_set_page_saves_everything_at_once_too(self):
        s = ConsumableSet.objects.create(name='Kit')
        a = ConsumableSetItem.objects.create(consumable_set=s, consumable=self.contrast, quantity=1, unit=self.ml)
        b = ConsumableSetItem.objects.create(consumable_set=s, consumable=self.cannula, quantity=1, unit=self.piece)
        url = reverse('set_save_all', args=[s.pk])
        def srow(i, q):
            return {f's{i.pk}-quantity': q, f's{i.pk}-unit': i.unit.pk, f's{i.pk}-is_required': 'on', f's{i.pk}-replaces': '', f's{i.pk}-notes': ''}
        self.client.post(url, {**srow(a, '7'), **srow(b, '8'), 'new0-consumable': self.glove.pk, 'new0-quantity': '2',
                               'new0-unit': self.piece.pk, 'new0-is_required': 'on'})
        a.refresh_from_db(); b.refresh_from_db()
        self.assertEqual((a.quantity, b.quantity, s.items.count()), (Decimal('7'), Decimal('8'), 3))
        r = self.client.post(url, {**srow(a, '1'), **srow(b, '2.5')})                       # invalid: nothing saved
        self.assertEqual(r.status_code, 200)
        a.refresh_from_db()
        self.assertEqual(a.quantity, Decimal('7'))
        page = self.client.get(reverse('set_detail', args=[s.pk])).content.decode()
        self.assertEqual(page.count('id="setform"'), 1)
        self.assertNotIn('/set-items/%d/update/' % a.pk, page)

    def test_set_page_permissions(self):
        s = ConsumableSet.objects.create(name='Kit2')
        a = ConsumableSetItem.objects.create(consumable_set=s, consumable=self.contrast, quantity=1, unit=self.ml)
        adder = client_for(make_user('adder11', edit=True))
        adder.post(reverse('set_save_all', args=[s.pk]), {f's{a.pk}-quantity': '9', f's{a.pk}-unit': self.ml.pk,
                                                          f's{a.pk}-is_required': 'on', f's{a.pk}-notes': ''})
        a.refresh_from_db()
        self.assertEqual(a.quantity, Decimal('1'))

    def test_old_single_row_endpoints_still_work_for_scripts(self):
        self.client.post(reverse('link_update', args=[self.l1.pk]), {
            f'l{self.l1.pk}-quantity': '6', f'l{self.l1.pk}-unit': self.ml.pk, f'l{self.l1.pk}-is_required': 'on',
            f'l{self.l1.pk}-notes': ''})
        self.l1.refresh_from_db()
        self.assertEqual(self.l1.quantity, Decimal('6'))



class SaveManyThingsAtOnce(BaseCase):
    """The single Save also takes several new items, optional/replacement links between them, and the
    page's other choices (apply a set, copy from an examination, save as a set)."""

    def setUp(self):
        super().setUp()
        self.url = reverse('exam_save_all', args=[self.exam.pk])
        self.pair = Unit.objects.get(name='pair')
        self.gloves = Consumable.objects.create(name='Gloves', default_unit=self.pair)
        self.gloves_pf = Consumable.objects.create(name='Gloves powder free', default_unit=self.pair)
        self.gloves_xl = Consumable.objects.create(name='Gloves XL', default_unit=self.pair)
        self.mask = Consumable.objects.create(name='Mask', default_unit=self.piece)

    def new(self, i, consumable, qty, unit, required=True, replaces='', notes=''):
        p = f'new{i}-'
        return {p + 'consumable': consumable.pk, p + 'quantity': qty, p + 'unit': unit.pk,
                p + 'is_required': 'on' if required else '', p + 'replaces': replaces, p + 'notes': notes}

    def test_required_item_and_its_replacements_in_one_click(self):
        data = {**self.new(0, self.gloves, 3, self.pair),
                **self.new(1, self.gloves_pf, 3, self.pair, required=False, replaces=self.gloves.pk, notes='allergy'),
                **self.new(2, self.gloves_xl, 3, self.pair, required=False, replaces=self.gloves.pk),
                **self.new(3, self.mask, 1, self.piece, required=False)}           # optional, replaces nothing
        r = self.client.post(self.url, data)
        self.assertEqual(r.status_code, 302)
        links = {l.consumable.name: l for l in self.exam.consumable_links.select_related('replaces', 'consumable')}
        self.assertEqual(len(links), 4)
        self.assertTrue(links['Gloves'].is_required)
        self.assertEqual((links['Gloves powder free'].is_required, links['Gloves powder free'].replaces,
                          links['Gloves powder free'].notes), (False, self.gloves, 'allergy'))
        self.assertEqual(links['Gloves XL'].replaces, self.gloves)
        self.assertIsNone(links['Mask'].replaces)
        msg = [str(m) for m in self.client.get(reverse('exam_detail', args=[self.exam.pk])).context['messages']]
        self.assertTrue(any(m.startswith('Saved: 4 added (') for m in msg), msg)

    def test_blank_rows_are_ignored_and_a_replacement_must_point_at_a_required_item(self):
        data = {**self.new(0, self.gloves_pf, 1, self.pair, required=False, replaces=self.mask.pk),   # mask not required
                'new1-consumable': '', 'new1-quantity': '1', 'new1-unit': '', 'new1-is_required': 'on'}
        r = self.client.post(self.url, data)
        self.assertEqual(r.status_code, 200)
        self.assertFalse(self.exam.consumable_links.exists())

    def test_the_same_consumable_twice_blocks_everything(self):
        data = {**self.new(0, self.gloves, 1, self.pair), **self.new(1, self.gloves, 2, self.pair)}
        r = self.client.post(self.url, data)
        self.assertEqual(r.status_code, 200)
        self.assertContains(r, 'listed twice')
        self.assertFalse(self.exam.consumable_links.exists())

    def test_one_bad_new_row_saves_nothing_and_all_typed_rows_come_back(self):
        data = {**self.new(0, self.gloves, 3, self.pair), **self.new(1, self.mask, '2.5', self.piece)}
        r = self.client.post(self.url, data)
        self.assertEqual(r.status_code, 200)
        self.assertFalse(self.exam.consumable_links.exists())
        html = r.content.decode()
        self.assertEqual(html.count('class="newrow"'), 3)               # the two typed rows + the hidden template
        self.assertIn('value="2.5"', html)
        self.assertIn('Use a whole number', html)

    def test_existing_row_can_become_a_replacement_for_an_item_added_in_the_same_click(self):
        old = ExaminationConsumable.objects.create(examination=self.exam, consumable=self.mask, quantity=1, unit=self.piece)
        data = {f'l{old.pk}-quantity': '1', f'l{old.pk}-unit': self.piece.pk, f'l{old.pk}-is_required': '',
                f'l{old.pk}-replaces': self.gloves.pk, f'l{old.pk}-notes': '',
                **self.new(0, self.gloves, 3, self.pair)}
        self.assertEqual(self.client.post(self.url, data).status_code, 302)
        old.refresh_from_db()
        self.assertEqual((old.is_required, old.replaces), (False, self.gloves))

    def test_apply_a_set_copy_and_save_as_set_with_the_same_button(self):
        kit = ConsumableSet.objects.create(name='Kit')
        ConsumableSetItem.objects.create(consumable_set=kit, consumable=self.gloves, quantity=2, unit=self.pair)
        ConsumableSetItem.objects.create(consumable_set=kit, consumable=self.gloves_pf, quantity=2, unit=self.pair,
                                         is_required=False, replaces=self.gloves)
        ExaminationConsumable.objects.create(examination=self.exam2, consumable=self.mask, quantity=4, unit=self.piece)
        data = {'apply_set': kit.pk, 'copy_from': self.exam2.pk, 'save_as_set': 'Everything here',
                **self.new(0, self.cannula, 1, self.piece)}
        self.assertEqual(self.client.post(self.url, data).status_code, 302)
        names = set(self.exam.consumable_links.values_list('consumable__name', flat=True))
        self.assertEqual(names, {'Gloves', 'Gloves powder free', 'Mask', 'IV Cannula'})
        self.assertEqual(self.exam.consumable_links.get(consumable=self.gloves_pf).replaces, self.gloves)   # link kept
        saved = ConsumableSet.objects.get(name='Everything here')
        self.assertEqual(saved.items.count(), 4)
        self.assertEqual(saved.items.get(consumable=self.gloves_pf).replaces, self.gloves)

    def test_extras_are_checked_before_anything_is_saved(self):
        ConsumableSet.objects.create(name='Taken')
        r = self.client.post(self.url, {'save_as_set': 'taken', **self.new(0, self.gloves, 1, self.pair)})
        self.assertEqual(r.status_code, 200)
        self.assertFalse(self.exam.consumable_links.exists())
        self.assertContains(r, 'already exists')
        r = self.client.post(self.url, {'save_as_set': 'Empty'})        # nothing to save as a set
        self.assertEqual(r.status_code, 200)
        self.assertFalse(ConsumableSet.objects.filter(name='Empty').exists())

    def test_page_has_rows_template_extras_and_one_form(self):
        html = self.client.get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        for needle in ('id="newitems"', 'id="add-new-row"', 'id="new-row-tpl"', 'rows.js', 'name="new0-consumable"',
                       'name="save_as_set"', 'new__i__'):
            self.assertIn(needle, html)
        self.assertEqual(html.count('<form id="examform"'), 1)
        self.assertNotIn('action="/examinations/%d/apply-set/"' % self.exam.pk, html)       # no separate forms any more
        self.assertNotIn('action="/examinations/%d/consumables/copy/"' % self.exam.pk, html)

    def test_sets_can_hold_optional_and_replacement_items_and_save_many_at_once(self):
        s = ConsumableSet.objects.create(name='Gloves kit')
        url = reverse('set_save_all', args=[s.pk])
        data = {**self.new(0, self.gloves, 3, self.pair), **self.new(1, self.gloves_pf, 3, self.pair, required=False,
                                                                     replaces=self.gloves.pk)}
        self.assertEqual(self.client.post(url, data).status_code, 302)
        self.assertEqual(s.items.get(consumable=self.gloves_pf).replaces, self.gloves)
        self.client.post(reverse('exam_apply_set', args=[self.exam.pk]), {'set': s.pk})
        self.assertEqual(self.exam.consumable_links.get(consumable=self.gloves_pf).replaces, self.gloves)
        page = self.client.get(reverse('set_detail', args=[s.pk])).content.decode()
        self.assertIn('name="new0-consumable"', page)
        self.assertIn('data-cid=', page)

    def test_applying_a_set_never_keeps_a_link_to_something_that_is_not_required_there(self):
        s = ConsumableSet.objects.create(name='Odd')
        ConsumableSetItem.objects.create(consumable_set=s, consumable=self.gloves_pf, quantity=1, unit=self.pair,
                                         is_required=False, replaces=self.gloves)
        from .services import apply_items, items_from_set
        apply_items([self.exam], items_from_set(s))
        self.assertIsNone(self.exam.consumable_links.get(consumable=self.gloves_pf).replaces)



class NothingBlocksSaveInTheBrowser(BaseCase):
    """Browsers silently refuse to submit a form that has an empty 'required' field. The blank 'add' row used to
    stop Save whenever only existing rows were edited."""

    def test_forms_skip_browser_validation_and_new_row_controls_are_not_required(self):
        import re
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast, quantity=1, unit=self.ml)
        html = self.client.get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        self.assertRegex(html, r'<form id="examform"[^>]* novalidate')
        for name in ('new0-consumable', 'new0-unit', 'new0-quantity'):
            tag = re.search(r'<(?:select|input)[^>]*name="%s"[^>]*>' % name, html).group(0)
            self.assertNotIn('required', tag, tag)
        s = ConsumableSet.objects.create(name='K')
        page = self.client.get(reverse('set_detail', args=[s.pk])).content.decode()
        self.assertRegex(page, r'<form id="setform"[^>]* novalidate')
        tag = re.search(r'<select[^>]*name="new0-consumable"[^>]*>', page).group(0)
        self.assertNotIn('required', tag)

    def test_only_editing_existing_rows_works_with_the_blank_row_left_alone(self):
        link = ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast, quantity=1, unit=self.ml)
        data = {f'l{link.pk}-quantity': '8', f'l{link.pk}-unit': self.ml.pk, f'l{link.pk}-is_required': 'on',
                f'l{link.pk}-replaces': '', f'l{link.pk}-notes': 'n',
                'new0-consumable': '', 'new0-quantity': '1', 'new0-unit': '', 'new0-is_required': 'on',
                'new0-replaces': '', 'new0-notes': ''}                       # exactly what a browser sends
        self.assertEqual(self.client.post(reverse('exam_save_all', args=[self.exam.pk]), data).status_code, 302)
        link.refresh_from_db()
        self.assertEqual((link.quantity, link.notes), (Decimal('8'), 'n'))


class QueuedRemoveAndSwap(BaseCase):
    """Remove and 'use instead' are saved by the same button, together with everything else."""

    def setUp(self):
        super().setUp()
        self.alt = Consumable.objects.create(name='Contrast B', default_unit=self.ml)
        ConsumableAlternative.objects.create(consumable=self.contrast, alternative=self.alt)
        self.a = ExaminationConsumable.objects.create(examination=self.exam, consumable=self.contrast, quantity=3, unit=self.ml)
        self.b = ExaminationConsumable.objects.create(examination=self.exam, consumable=self.cannula, quantity=1, unit=self.piece)
        self.opt = Consumable.objects.create(name='Mask', default_unit=self.piece)
        self.c = ExaminationConsumable.objects.create(examination=self.exam, consumable=self.opt, quantity=1, unit=self.piece,
                                                      is_required=False, replaces=self.cannula)
        self.url = reverse('exam_save_all', args=[self.exam.pk])

    def row(self, link, qty=None, **extra):
        p = f'l{link.pk}-'
        return {p + 'quantity': qty or link.quantity, p + 'unit': link.unit.pk,
                p + 'is_required': 'on' if link.is_required else '', p + 'replaces': link.replaces_id or '',
                p + 'notes': '', **{p + k: v for k, v in extra.items()}}

    def test_everything_in_one_save(self):
        gloves = Consumable.objects.create(name='Gloves', default_unit=self.piece)
        data = {**self.row(self.a, swap=self.alt.pk, quantity='4'),                  # edit AND swap
                **self.row(self.b, remove='1'),                                      # remove (its optional link is cleared)
                **self.row(self.c, qty='2'),
                'new0-consumable': gloves.pk, 'new0-quantity': '1', 'new0-unit': self.piece.pk, 'new0-is_required': 'on'}
        r = self.client.post(self.url, data, follow=True)
        names = set(self.exam.consumable_links.values_list('consumable__name', flat=True))
        self.assertEqual(names, {'Contrast B', 'Mask', 'Gloves'})
        self.a.refresh_from_db(); self.c.refresh_from_db()
        self.assertEqual((self.a.consumable, self.a.quantity), (self.alt, Decimal('4')))
        self.assertEqual((self.c.quantity, self.c.replaces), (Decimal('2'), None))     # its target was removed
        msg = ' '.join(str(m) for m in r.context['messages'])
        self.assertIn('1 swapped for an alternative', msg)
        self.assertIn('1 removed', msg)

    def test_removing_needs_the_delete_permission_and_refuses_the_whole_save(self):
        changer = client_for(make_user('chg', edit=True, change=True))
        changer.post(self.url, {**self.row(self.a, qty='9'), **self.row(self.b, remove='1')})
        self.a.refresh_from_db()
        self.assertEqual(self.a.quantity, Decimal('3'))
        self.assertTrue(ExaminationConsumable.objects.filter(pk=self.b.pk).exists())
        deleter = client_for(make_user('del', edit=True, change=True, delete=True))
        deleter.post(self.url, {**self.row(self.a, qty='9'), **self.row(self.b, remove='1')})
        self.a.refresh_from_db()
        self.assertEqual(self.a.quantity, Decimal('9'))
        self.assertFalse(ExaminationConsumable.objects.filter(pk=self.b.pk).exists())

    def test_swap_needs_the_change_permission(self):
        adder = client_for(make_user('addr', edit=True))
        adder.post(self.url, {f'l{self.a.pk}-swap': self.alt.pk})
        self.a.refresh_from_db()
        self.assertEqual(self.a.consumable, self.contrast)

    def test_bad_swaps_save_nothing(self):
        r = self.client.post(self.url, {**self.row(self.a, qty='9'), **self.row(self.b, swap=self.alt.pk)})   # not b's alternative
        self.assertEqual(r.status_code, 200)
        self.a.refresh_from_db()
        self.assertEqual(self.a.quantity, Decimal('3'))
        ExaminationConsumable.objects.create(examination=self.exam, consumable=self.alt, quantity=1, unit=self.ml, is_required=False)
        r = self.client.post(self.url, {**self.row(self.a, swap=self.alt.pk)})                # already on the examination
        self.assertContains(r, 'already on this examination')

    def test_the_page_shows_ticks_and_a_choice_instead_of_separate_buttons(self):
        html = self.client.get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        self.assertIn('name="l%d-remove"' % self.a.pk, html)
        self.assertIn('name="l%d-swap"' % self.a.pk, html)
        self.assertNotIn('/exam-consumables/%d/remove/' % self.a.pk, html)
        self.assertNotIn('/exam-consumables/%d/swap/' % self.a.pk, html)
        editor = client_for(make_user('ed2', edit=True, change=True)).get(reverse('exam_detail', args=[self.exam.pk])).content.decode()
        self.assertNotIn('-remove"', editor)
        self.assertIn('-swap"', editor)

    def test_failed_save_redisplays_ticks_and_swap_choices(self):
        r = self.client.post(self.url, {**self.row(self.a, swap=self.alt.pk, qty='2.5'), **self.row(self.b, remove='1')})
        html = r.content.decode()
        self.assertEqual(r.status_code, 200)
        self.assertRegex(html, r'name="l%d-remove"[^>]* checked' % self.b.pk)
        self.assertRegex(html, r'<option value="%d" selected>Use Contrast B instead' % self.alt.pk)
        self.assertTrue(ExaminationConsumable.objects.filter(pk=self.b.pk).exists())

    def test_sets_remove_on_save_too(self):
        s = ConsumableSet.objects.create(name='K')
        i1 = ConsumableSetItem.objects.create(consumable_set=s, consumable=self.contrast, quantity=1, unit=self.ml)
        i2 = ConsumableSetItem.objects.create(consumable_set=s, consumable=self.cannula, quantity=1, unit=self.piece)
        def srow(i, **x):
            return {f's{i.pk}-quantity': 5, f's{i.pk}-unit': i.unit.pk, f's{i.pk}-is_required': 'on',
                    f's{i.pk}-replaces': '', f's{i.pk}-notes': '', **{f's{i.pk}-{k}': v for k, v in x.items()}}
        self.client.post(reverse('set_save_all', args=[s.pk]), {**srow(i1), **srow(i2, remove='1')})
        self.assertEqual(list(s.items.values_list('consumable__name', flat=True)), ['Contrast'])
        i1.refresh_from_db()
        self.assertEqual(i1.quantity, Decimal('5'))
        html = self.client.get(reverse('set_detail', args=[s.pk])).content.decode()
        self.assertIn('name="s%d-remove"' % i1.pk, html)
