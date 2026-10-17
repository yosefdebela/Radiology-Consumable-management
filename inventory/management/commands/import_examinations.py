from decimal import Decimal, InvalidOperation
from pathlib import Path

from django.conf import settings
from django.core.management.base import BaseCommand, CommandError
from openpyxl import load_workbook

from inventory.models import Examination, Modality

# sheet name -> modality slug
SHEETS = {'CT': 'ct', 'MRI': 'mri', 'Ultrasound': 'ultrasound',
          'X-Ray': 'x-ray', 'Special X-Ray': 'special-x-ray'}


def clean_price(v):
    if v in (None, ''):
        return None
    try:
        return Decimal(str(v))
    except InvalidOperation:
        return None


class Command(BaseCommand):
    help = 'Import examinations from the categorized Excel workbook (safe to re-run; matches on code).'

    def add_arguments(self, parser):
        parser.add_argument('path', nargs='?',
                            default=str(Path(settings.BASE_DIR) / 'data' / 'Radiology_Investigations_Categorized.xlsx'))

    def handle(self, *args, **opts):
        path = Path(opts['path'])
        if not path.exists():
            raise CommandError(f'File not found: {path}')
        wb = load_workbook(path, read_only=True, data_only=True)
        created = updated = 0
        for sheet, slug in SHEETS.items():
            if sheet not in wb.sheetnames:
                self.stderr.write(f'Sheet "{sheet}" not found, skipped.')
                continue
            modality = Modality.objects.get(slug=slug)
            rows = wb[sheet].iter_rows(values_only=True)
            header = [str(h).strip().lower() if h else '' for h in next(rows)]
            idx = {h: i for i, h in enumerate(header)}
            for r in rows:
                name = (r[idx['examination']] or '').strip() if r[idx['examination']] else ''
                if not name:
                    continue
                code = str(r[idx['code']] or '').strip()
                defaults = {
                    'name': ' '.join(name.split()),
                    'modality': modality,
                    'subcategory': (r[idx['subcategory']] or '').strip() if 'subcategory' in idx and r[idx['subcategory']] else '',
                    'price': clean_price(r[idx['price']]) if 'price' in idx else None,
                }
                obj, was_created = Examination.objects.update_or_create(code=code, defaults=defaults) \
                    if code else Examination.objects.get_or_create(name=defaults['name'], modality=modality,
                                                                   defaults=defaults)
                created += was_created
                updated += (not was_created)
        self.stdout.write(self.style.SUCCESS(f'Examinations imported: {created} new, {updated} updated.'))
