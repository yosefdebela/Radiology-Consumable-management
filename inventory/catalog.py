"""Import the consumables list (Excel) used for auto-fill suggestions."""
import re
from io import BytesIO

from django.db import transaction
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import CatalogItem, Consumable, ConsumableCategory, Unit, catalog_key
from .services import set_alternatives, split_names

HEADERS = {
    'name': {'name', 'consumable', 'consumable name', 'consumables', 'item', 'item name', 'material',
             'material name', 'product', 'product name', 'description of item'},
    'category': {'category', 'consumable category', 'type', 'group', 'class', 'item category'},
    'unit': {'unit', 'units', 'uom', 'unit of measure', 'unit of measurement', 'measurement unit', 'measurement'},
    'alternatives': {'alternative', 'alternatives', 'alternative item', 'alternative items', 'substitute',
                     'substitutes', 'replacement', 'replacements'},
    'description': {'description', 'notes', 'note', 'details', 'specification', 'remarks', 'comment'},
}
UNIT_ALIASES = {'pcs': 'piece', 'pc': 'piece', 'pieces': 'piece', 'each': 'piece', 'prs': 'pair', 'pairs': 'pair',
                'btl': 'bottle', 'bottles': 'bottle', 'packs': 'pack', 'pkt': 'pack', 'packet': 'pack',
                'vials': 'vial', 'mls': 'ml', 'milliliter': 'ml', 'millilitre': 'ml', 'boxes': 'box',
                'rolls': 'roll', 'sets': 'set'}
MAX_NAME, MAX_CATEGORY, MAX_DESC = 150, 80, 255


def _norm_header(h):
    return re.sub(r'[\s_\-/.]+', ' ', str(h or '').strip().lower()).strip()


def _text(v):
    return '' if v is None else ' '.join(str(v).split())


def _map_columns(header):
    cols = {}
    for i, h in enumerate(header):
        n = _norm_header(h)
        for field, aliases in HEADERS.items():
            if n in aliases and field not in cols:
                cols[field] = i
                break
    return cols


def import_catalog(source, add_as_consumables=False):
    """Read an .xlsx (path or file object). Updates existing entries, never deletes.

    Returns a report dict. Categories and units that do not exist yet are created.
    """
    wb = load_workbook(source, read_only=True, data_only=True)
    rep = {'new': 0, 'updated': 0, 'skipped': 0, 'categories_created': 0, 'units_created': [],
           'consumables_added': 0, 'sheets_used': [], 'sheets_skipped': []}
    cat_cache = {c.name.lower(): c for c in ConsumableCategory.objects.all()}
    unit_cache = {u.name.lower(): u for u in Unit.objects.all()}

    def category(name):
        if not name:
            return None
        name = name[:MAX_CATEGORY]
        obj = cat_cache.get(name.lower())
        if obj is None:
            obj = cat_cache[name.lower()] = ConsumableCategory.objects.create(name=name)
            rep['categories_created'] += 1
        return obj

    def unit(name):
        if not name:
            return None
        key = UNIT_ALIASES.get(name.lower(), name.lower())[:30]
        obj = unit_cache.get(key)
        if obj is None:
            obj = unit_cache[key] = Unit.objects.create(name=key)
            rep['units_created'].append(key)
        return obj

    seen_keys = []
    with transaction.atomic():
        single = len(wb.sheetnames) == 1
        for ws in wb.worksheets:
            rows = ws.iter_rows(values_only=True)
            header = next(rows, None)
            if not header:
                rep['sheets_skipped'].append(ws.title)
                continue
            cols = _map_columns(header)
            if 'name' not in cols:
                if not single:
                    rep['sheets_skipped'].append(ws.title)
                    continue
                cols['name'] = 0          # a single sheet: assume the first column holds the names
            rep['sheets_used'].append(ws.title)

            def cell(row, field):
                i = cols.get(field)
                return _text(row[i]) if i is not None and i < len(row) else ''

            for row in rows:
                name = cell(row, 'name')
                if not any(_text(c) for c in row):
                    continue
                if not name or len(name) > MAX_NAME:
                    rep['skipped'] += 1
                    continue
                item, created = CatalogItem.objects.get_or_create(name_key=catalog_key(name),
                                                                  defaults={'name': name})
                cat, un, desc = category(cell(row, 'category')), unit(cell(row, 'unit')), cell(row, 'description')[:MAX_DESC]
                alts = '; '.join(split_names(cell(row, 'alternatives')))[:500]
                changed = created
                item.name = name
                if cat and item.category_id != cat.pk:
                    item.category, changed = cat, True
                if un and item.default_unit_id != un.pk:
                    item.default_unit, changed = un, True
                if desc and item.description != desc:
                    item.description, changed = desc, True
                if alts and item.alternatives != alts:
                    item.alternatives, changed = alts, True
                item.save()
                if created:
                    rep['new'] += 1
                elif changed:
                    rep['updated'] += 1
                seen_keys.append(item.name_key)

        if add_as_consumables and seen_keys:
            fallback = None
            for item in CatalogItem.objects.filter(name_key__in=seen_keys).select_related('category', 'default_unit'):
                # checked per item: an earlier item's alternatives may already have created this consumable
                if Consumable.objects.filter(name__iexact=item.name).exists():
                    continue
                if item.default_unit is None and fallback is None:
                    fallback = unit('piece')
                made = Consumable.objects.create(name=item.name, category=item.category,
                                                 default_unit=item.default_unit or fallback,
                                                 description=item.description)
                rep['consumables_added'] += 1
                if item.alternatives:
                    set_alternatives(made, split_names(item.alternatives))
    return rep


def build_template():
    wb = Workbook()
    ws = wb.active
    ws.title = 'Consumables'
    ws.append(['Name', 'Category', 'Unit', 'Description', 'Alternatives'])
    for c in ws[1]:
        c.font, c.fill = Font(bold=True, color='FFFFFF'), PatternFill('solid', fgColor='12262B')
    for row in [['Iodinated contrast (CT)', 'Contrast media', 'ml', 'Non-ionic, for intravenous use',
                 'Iodinated contrast 370; Contrast 300'],
                ['IV cannula 20G', 'IV and injection supplies', 'piece', '', 'IV cannula 22G'],
                ['Gloves', 'Gloves and protection', 'pair', 'Sterile', '']]:
        ws.append(row)
    for i, w in enumerate([36, 28, 12, 44, 40], start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
